import sqlite3
import os
import shutil
import urllib.request
import urllib.parse
import webbrowser
from datetime import datetime

# ===== مكتبات PDF =====
try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.pdfgen import canvas
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False
    print("تحذير: reportlab غير مثبت.")

from kivy.app import App
from kivy.core.text import LabelBase
from kivy.uix.boxlayout import BoxLayout
from kivy.uix.gridlayout import GridLayout
from kivy.uix.label import Label
from kivy.uix.textinput import TextInput
from kivy.uix.button import Button
from kivy.uix.image import Image
from kivy.uix.popup import Popup
from kivy.uix.screenmanager import ScreenManager, Screen
from kivy.uix.scrollview import ScrollView
from kivy.uix.relativelayout import RelativeLayout
from kivy.uix.spinner import Spinner, SpinnerOption
from kivy.graphics import Color, RoundedRectangle, Rectangle
import openpyxl

import arabic_reshaper
from bidi.algorithm import get_display

# ============================================================
# 1. تهيئة الخطوط
# ============================================================
FONT_NAME = "AmiriFont"
FONT_FILE = "Amiri-Regular.ttf"
FONT_BOLD_NAME = "AmiriBold"
FONT_BOLD_FILE = "Amiri-Bold.ttf"

if not os.path.exists(FONT_FILE):
    try:
        url = "https://github.com/google/fonts/raw/main/ofl/amiri/Amiri-Regular.ttf"
        urllib.request.urlretrieve(url, FONT_FILE)
    except Exception as e:
        print(f"Amiri font download info: {e}")

if os.path.exists(FONT_FILE):
    LabelBase.register(name=FONT_NAME, fn_regular=FONT_FILE)
    ARABIC_FONT = FONT_NAME
else:
    ARABIC_FONT = None

if not os.path.exists(FONT_BOLD_FILE):
    try:
        url = "https://github.com/google/fonts/raw/main/ofl/amiri/Amiri-Bold.ttf"
        urllib.request.urlretrieve(url, FONT_BOLD_FILE)
    except Exception as e:
        print(f"Amiri Bold font download info: {e}")

if os.path.exists(FONT_BOLD_FILE):
    LabelBase.register(name=FONT_BOLD_NAME, fn_regular=FONT_BOLD_FILE)
    ARABIC_FONT_BOLD = FONT_BOLD_NAME
else:
    ARABIC_FONT_BOLD = ARABIC_FONT

if PDF_AVAILABLE and os.path.exists(FONT_FILE):
    try:
        pdfmetrics.registerFont(TTFont('Amiri', FONT_FILE))
        if os.path.exists(FONT_BOLD_FILE):
            pdfmetrics.registerFont(TTFont('Amiri-Bold', FONT_BOLD_FILE))
    except Exception as e:
        print(f"PDF font registration error: {e}")


def ar(text):
    if not text:
        return ""
    reshaped_text = arabic_reshaper.reshape(str(text))
    return get_display(reshaped_text)


def fmt_num(n):
    try:
        return f"{float(n):,.2f}"
    except (ValueError, TypeError):
        return "0.00"


class ArabicSpinnerOption(SpinnerOption):
    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        if ARABIC_FONT:
            self.font_name = ARABIC_FONT
        self.font_size = '16sp'


def get_logo_path():
    possible_names = [
        "logo.png", "logo.jpg", "logo.jpeg",
        "logo_4.jpg", "logo_3.jpg", "logo_2.jpg",
        "شعار.png", "شعار.jpg"
    ]
    for name in possible_names:
        if os.path.exists(name):
            return name
    for f in os.listdir('.'):
        if f.lower().endswith(('.jpg', '.jpeg', '.png')) and ('logo' in f.lower() or 'شعار' in f):
            return f
    return None


# ============================================================
# 2. عنصر إدخال مخصص
# ============================================================
class ArabicTextInput(RelativeLayout):
    def __init__(self, text='', hint_text='', multiline=False, input_filter=None, password=False, **kwargs):
        super().__init__(**kwargs)
        lbl_kwargs = {'font_name': ARABIC_FONT} if ARABIC_FONT else {}

        self.real_input = TextInput(
            text=str(text) if text else '',
            multiline=multiline,
            password=password,
            input_filter=input_filter,
            cursor_color=(0.2, 0.6, 1, 1),
            foreground_color=(0, 0, 0, 0),
            background_color=(0.12, 0.15, 0.22, 1),
            padding=(12, 12),
            size_hint=(1, 1),
            pos_hint={'x': 0, 'y': 0}
        )

        display_txt = "*" * len(str(text)) if password and text else (str(text) if text else hint_text)
        self.display_label = Label(
            text=ar(display_txt),
            font_size='14sp',
            color=(1, 1, 1, 1) if text else (0.55, 0.6, 0.7, 1),
            halign='right',
            valign='middle',
            size_hint=(1, 1),
            pos_hint={'x': 0, 'y': 0},
            **lbl_kwargs
        )
        self.display_label.bind(size=self._update_text_size)
        self.hint_text = hint_text
        self.password = password
        self.real_input.bind(text=self.on_text_change)

        self.add_widget(self.real_input)
        self.add_widget(self.display_label)

    def _update_text_size(self, instance, value):
        self.display_label.text_size = (value[0] - 20, value[1])

    def on_text_change(self, instance, value):
        if value:
            txt = "*" * len(value) if self.password else value
            self.display_label.text = ar(txt)
            self.display_label.color = (1, 1, 1, 1)
        else:
            self.display_label.text = ar(self.hint_text)
            self.display_label.color = (0.55, 0.6, 0.7, 1)

    @property
    def text(self):
        return self.real_input.text

    @text.setter
    def text(self, val):
        self.real_input.text = str(val) if val is not None else ''


# ============================================================
# 3. قاعدة البيانات
# ============================================================
DB_NAME = "customs_erp_v22.db"
BACKUP_DIR = "backups"
INVOICES_DIR = "invoices"


def get_conn():
    return sqlite3.connect(DB_NAME)


def init_db():
    conn = get_conn()
    cursor = conn.cursor()

    cursor.execute('''
        CREATE TABLE IF NOT EXISTS users (
            username TEXT PRIMARY KEY,
            password TEXT,
            full_name TEXT
        )
    ''')
    cursor.execute("INSERT OR IGNORE INTO users VALUES ('admin', '1234', 'أدهم الحاج')")

    cursor.execute('''
        CREATE TABLE IF NOT EXISTS clients (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            full_name TEXT NOT NULL UNIQUE,
            phone TEXT,
            email TEXT,
            address TEXT,
            tax_number TEXT,
            credit_limit REAL DEFAULT 0,
            notes TEXT,
            created_at TEXT
        )
    ''')

    cursor.execute('''
        CREATE TABLE IF NOT EXISTS exp_categories (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            cat_name TEXT UNIQUE
        )
    ''')
    default_cats = [('تشغيلي',), ('إداري',), ('إيجارات',), ('مواصلات ونقل',),
                    ('صيانة ومعدات',), ('نثريات وضيافة',)]
    cursor.executemany("INSERT OR IGNORE INTO exp_categories (cat_name) VALUES (?)", default_cats)

    cursor.execute('''
        CREATE TABLE IF NOT EXISTS fee_types (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            fee_name TEXT UNIQUE
        )
    ''')
    default_fees = [('الرسوم الجمركية الأساسية',), ('رسوم الأرضيات والغرامات',),
                    ('مصاريف النقل والتحسين',), ('رسوم الجودة',), ('رسوم الزراعة',),
                    ('رسوم التأمين',), ('أتعاب التخليص الجمركي',)]
    cursor.executemany("INSERT OR IGNORE INTO fee_types (fee_name) VALUES (?)", default_fees)

    cursor.execute('''
        CREATE TABLE IF NOT EXISTS shipments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            shipment_no TEXT UNIQUE,
            client_id INTEGER,
            client_name TEXT,
            currency TEXT DEFAULT 'YER',
            advance_amount REAL,
            customs_duty REAL,
            storage_fee REAL,
            transport_fee REAL,
            quality_fee REAL,
            agriculture_fee REAL,
            insurance_fee REAL,
            other_fee_amount REAL DEFAULT 0,
            clearance_fee REAL,
            total_expenses REAL,
            balance REAL,
            doc_path TEXT DEFAULT '',
            status TEXT DEFAULT 'تحت الإجراء',
            created_at TEXT
        )
    ''')

    cursor.execute('''
        CREATE TABLE IF NOT EXISTS shipment_custom_fees (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            shipment_no TEXT,
            fee_name TEXT,
            amount REAL
        )
    ''')

    cursor.execute('''
        CREATE TABLE IF NOT EXISTS journal_entries (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            entry_date TEXT,
            ref_no TEXT,
            account_code TEXT,
            debit REAL DEFAULT 0,
            credit REAL DEFAULT 0,
            currency TEXT DEFAULT 'YER',
            description TEXT
        )
    ''')

    cursor.execute('''
        CREATE TABLE IF NOT EXISTS expenses (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            exp_date TEXT,
            exp_title TEXT,
            category TEXT,
            amount REAL,
            currency TEXT DEFAULT 'YER'
        )
    ''')

    cursor.execute('''
        CREATE TABLE IF NOT EXISTS invoices (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            invoice_no TEXT UNIQUE,
            shipment_no TEXT,
            client_id INTEGER,
            client_name TEXT,
            amount REAL,
            currency TEXT DEFAULT 'YER',
            pdf_path TEXT,
            notes TEXT,
            created_at TEXT
        )
    ''')

    conn.commit()
    conn.close()

    migrate_existing_clients()


def migrate_existing_clients():
    conn = get_conn()
    cursor = conn.cursor()

    cursor.execute("""
        SELECT DISTINCT client_name FROM shipments
        WHERE client_id IS NULL AND client_name IS NOT NULL AND client_name != ''
    """)
    names = [r[0] for r in cursor.fetchall()]

    for name in names:
        cursor.execute("""
            INSERT OR IGNORE INTO clients (full_name, created_at)
            VALUES (?, ?)
        """, (name, datetime.now().strftime("%Y-%m-%d %H:%M")))

        cursor.execute("SELECT id FROM clients WHERE full_name = ?", (name,))
        row = cursor.fetchone()
        if row:
            cursor.execute("""
                UPDATE shipments SET client_id = ?
                WHERE client_name = ? AND client_id IS NULL
            """, (row[0], name))

    conn.commit()
    conn.close()


init_db()


def backup_database():
    try:
        if not os.path.exists(BACKUP_DIR):
            os.makedirs(BACKUP_DIR)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup_file = os.path.join(BACKUP_DIR, f"backup_{timestamp}.db")
        shutil.copy2(DB_NAME, backup_file)
        return backup_file
    except Exception as e:
        return f"فشل النسخ: {str(e)}"


def post_shipment_accounting(shp_no, client_name, curr, adv, duty, storage, transport,
                             quality, agri, ins, other_amt, fee):
    conn = get_conn()
    cursor = conn.cursor()
    today = datetime.now().strftime("%Y-%m-%d %H:%M")

    cursor.execute("DELETE FROM journal_entries WHERE ref_no=?", (shp_no,))

    if adv > 0:
        cursor.execute(
            "INSERT INTO journal_entries (entry_date, ref_no, account_code, debit, credit, currency, description) VALUES (?,?,?,?,?,?,?)",
            (today, shp_no, '1101', adv, 0, curr, f"قبض عهدة شحنة {shp_no} - {client_name}"))
        cursor.execute(
            "INSERT INTO journal_entries (entry_date, ref_no, account_code, debit, credit, currency, description) VALUES (?,?,?,?,?,?,?)",
            (today, shp_no, '2101', 0, adv, curr, f"قيد أمانة عهدة شحنة {shp_no} - {client_name}"))

    pass_through_exp = duty + storage + transport + quality + agri + ins + other_amt
    if pass_through_exp > 0:
        cursor.execute(
            "INSERT INTO journal_entries (entry_date, ref_no, account_code, debit, credit, currency, description) VALUES (?,?,?,?,?,?,?)",
            (today, shp_no, '2101', pass_through_exp, 0, curr, f"رسوم ومصاريف شحنة {shp_no}"))
        cursor.execute(
            "INSERT INTO journal_entries (entry_date, ref_no, account_code, debit, credit, currency, description) VALUES (?,?,?,?,?,?,?)",
            (today, shp_no, '1101', 0, pass_through_exp, curr, f"سداد رسوم ومصاريف شحنة {shp_no}"))

    if fee > 0:
        cursor.execute(
            "INSERT INTO journal_entries (entry_date, ref_no, account_code, debit, credit, currency, description) VALUES (?,?,?,?,?,?,?)",
            (today, shp_no, '2101', fee, 0, curr, f"استحقاق أتعاب تخليص شحنة {shp_no}"))
        cursor.execute(
            "INSERT INTO journal_entries (entry_date, ref_no, account_code, debit, credit, currency, description) VALUES (?,?,?,?,?,?,?)",
            (today, shp_no, '4101', 0, fee, curr, f"إثبات إيراد أتعاب شحنة {shp_no}"))

    conn.commit()
    conn.close()# ============================================================
# 4. دوال العملاء (CRM)
# ============================================================
def get_all_clients_objects():
    conn = get_conn()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT id, full_name, phone, email, address, tax_number,
               credit_limit, notes, created_at
        FROM clients ORDER BY full_name
    """)
    rows = cursor.fetchall()
    conn.close()
    return rows


def get_client_by_id(client_id):
    conn = get_conn()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT id, full_name, phone, email, address, tax_number,
               credit_limit, notes, created_at
        FROM clients WHERE id = ?
    """, (client_id,))
    row = cursor.fetchone()
    conn.close()
    return row


def save_client(client_id, full_name, phone, email, address, tax_number, credit_limit, notes):
    conn = get_conn()
    cursor = conn.cursor()
    try:
        if client_id:
            cursor.execute("""
                UPDATE clients SET
                    full_name = ?, phone = ?, email = ?, address = ?,
                    tax_number = ?, credit_limit = ?, notes = ?
                WHERE id = ?
            """, (full_name, phone, email, address, tax_number, credit_limit, notes, client_id))
        else:
            cursor.execute("""
                INSERT INTO clients
                (full_name, phone, email, address, tax_number, credit_limit, notes, created_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """, (full_name, phone, email, address, tax_number, credit_limit, notes,
                  datetime.now().strftime("%Y-%m-%d %H:%M")))
        conn.commit()
        return True, "تم الحفظ بنجاح"
    except sqlite3.IntegrityError:
        return False, "اسم العميل موجود مسبقاً"
    except Exception as e:
        return False, f"خطأ: {str(e)}"
    finally:
        conn.close()


def delete_client(client_id):
    conn = get_conn()
    cursor = conn.cursor()
    cursor.execute("SELECT COUNT(*) FROM shipments WHERE client_id = ?", (client_id,))
    count = cursor.fetchone()[0]

    if count > 0:
        conn.close()
        return False, f"لا يمكن حذف العميل — لديه {count} شحنة مسجلة"

    try:
        cursor.execute("DELETE FROM clients WHERE id = ?", (client_id,))
        conn.commit()
        conn.close()
        return True, "تم حذف العميل"
    except Exception as e:
        conn.close()
        return False, f"خطأ: {str(e)}"


def get_client_stats(client_id):
    conn = get_conn()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT COUNT(*),
               COALESCE(SUM(advance_amount), 0),
               COALESCE(SUM(total_expenses), 0),
               COALESCE(SUM(balance), 0)
        FROM shipments WHERE client_id = ?
    """, (client_id,))
    row = cursor.fetchone()
    conn.close()
    return {
        'count': row[0],
        'total_adv': row[1],
        'total_exp': row[2],
        'balance': row[3],
    }


def get_client_statement(client_name, date_from=None, date_to=None):
    conn = get_conn()
    cursor = conn.cursor()
    query = """
        SELECT shipment_no, currency, advance_amount, total_expenses, clearance_fee,
               balance, status, created_at
        FROM shipments WHERE client_name = ?
    """
    params = [client_name]
    if date_from:
        query += " AND created_at >= ?"
        params.append(date_from)
    if date_to:
        query += " AND created_at <= ?"
        params.append(date_to)
    query += " ORDER BY id DESC"
    cursor.execute(query, tuple(params))
    rows = cursor.fetchall()
    conn.close()
    return rows


def get_all_clients():
    conn = get_conn()
    cursor = conn.cursor()
    cursor.execute("SELECT DISTINCT client_name FROM shipments ORDER BY client_name")
    rows = [r[0] for r in cursor.fetchall() if r[0]]
    conn.close()
    return rows


# ============================================================
# 5. دوال الفواتير
# ============================================================
def generate_invoice_no():
    conn = get_conn()
    cursor = conn.cursor()
    year = datetime.now().year

    cursor.execute("""
        SELECT invoice_no FROM invoices
        WHERE invoice_no LIKE ?
        ORDER BY id DESC LIMIT 1
    """, (f"INV-{year}-%",))
    row = cursor.fetchone()

    if row and row[0]:
        try:
            last_num = int(row[0].split('-')[-1])
            next_num = last_num + 1
        except ValueError:
            next_num = 1
    else:
        next_num = 1

    conn.close()
    return f"INV-{year}-{next_num:04d}"


def save_invoice(invoice_no, shipment_no, client_id, client_name, amount, currency, pdf_path, notes=""):
    conn = get_conn()
    cursor = conn.cursor()
    try:
        cursor.execute("""
            INSERT INTO invoices
            (invoice_no, shipment_no, client_id, client_name, amount, currency, pdf_path, notes, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (invoice_no, shipment_no, client_id, client_name, amount, currency, pdf_path, notes,
              datetime.now().strftime("%Y-%m-%d %H:%M")))
        conn.commit()
        return True
    except Exception as e:
        print(f"save_invoice error: {e}")
        return False
    finally:
        conn.close()


def get_all_invoices():
    conn = get_conn()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT id, invoice_no, shipment_no, client_name, amount, currency,
               pdf_path, notes, created_at
        FROM invoices ORDER BY id DESC
    """)
    rows = cursor.fetchall()
    conn.close()
    return rows


def delete_invoice_by_id(inv_id):
    conn = get_conn()
    cursor = conn.cursor()
    cursor.execute("SELECT pdf_path FROM invoices WHERE id = ?", (inv_id,))
    row = cursor.fetchone()
    if row and row[0] and os.path.exists(row[0]):
        try:
            os.remove(row[0])
        except Exception:
            pass

    cursor.execute("DELETE FROM invoices WHERE id = ?", (inv_id,))
    conn.commit()
    conn.close()


def generate_invoice_pdf(shipment_no, invoice_no=None):
    if not PDF_AVAILABLE:
        return None, "مكتبة reportlab غير مثبتة"

    conn = get_conn()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT shipment_no, client_id, client_name, currency, advance_amount,
               customs_duty, storage_fee, transport_fee, quality_fee, agriculture_fee,
               insurance_fee, other_fee_amount, clearance_fee, total_expenses, balance,
               created_at
        FROM shipments WHERE shipment_no = ?
    """, (shipment_no,))
    row = cursor.fetchone()

    if not row:
        conn.close()
        return None, "الشحنة غير موجودة"

    cursor.execute("""
        SELECT fee_name, amount FROM shipment_custom_fees WHERE shipment_no = ?
    """, (shipment_no,))
    custom_fees = cursor.fetchall()

    client_info = None
    if row[1]:
        cursor.execute("""
            SELECT full_name, phone, email, address, tax_number
            FROM clients WHERE id = ?
        """, (row[1],))
        client_info = cursor.fetchone()

    conn.close()

    if not os.path.exists(INVOICES_DIR):
        os.makedirs(INVOICES_DIR)

    if not invoice_no:
        invoice_no = generate_invoice_no()

    filename = os.path.join(INVOICES_DIR, f"{invoice_no}.pdf")

    try:
        c = canvas.Canvas(filename, pagesize=A4)
        width, height = A4

        logo_path = get_logo_path()
        if logo_path and os.path.exists(logo_path):
            try:
                c.drawImage(logo_path, width - 180, height - 140, width=140, height=100,
                            preserveAspectRatio=True, mask='auto')
            except Exception:
                pass

        c.setFont('Amiri-Bold' if os.path.exists(FONT_BOLD_FILE) else 'Amiri', 16)
        c.setFillColor(colors.HexColor('#0B4F9E'))
        c.drawString(20*mm, height - 30*mm, ar("مكتب أدهم الحاج"))
        c.setFont('Amiri', 11)
        c.setFillColor(colors.HexColor('#333333'))
        c.drawString(20*mm, height - 38*mm, ar("للتخليص الجمركي والخدمات اللوجستية"))
        c.drawString(20*mm, height - 45*mm, ar("جميع المنافذ الجمركية"))
        c.drawString(20*mm, height - 52*mm, ar("الجمهورية اليمنية"))

        c.setStrokeColor(colors.HexColor('#0B4F9E'))
        c.setLineWidth(2)
        c.line(15*mm, height - 60*mm, width - 15*mm, height - 60*mm)

        c.setFont('Amiri-Bold' if os.path.exists(FONT_BOLD_FILE) else 'Amiri', 22)
        c.setFillColor(colors.HexColor('#0B4F9E'))
        c.drawCentredString(width / 2, height - 75*mm, ar("فاتورة رسمية"))

        y = height - 90*mm

        c.setFont('Amiri-Bold' if os.path.exists(FONT_BOLD_FILE) else 'Amiri', 11)
        c.setFillColor(colors.black)

        c.drawString(20*mm, y, ar(f"رقم الفاتورة: {invoice_no}"))
        c.drawRightString(width - 20*mm, y, ar(f"التاريخ: {datetime.now().strftime('%Y-%m-%d')}"))
        y -= 7*mm

        c.drawString(20*mm, y, ar(f"رقم الشحنة: {row[0]}"))
        c.drawRightString(width - 20*mm, y, ar(f"العملة: {row[3]}"))
        y -= 10*mm

        c.setFillColor(colors.HexColor('#F0F4F8'))
        c.rect(15*mm, y - 30*mm, width - 30*mm, 32*mm, fill=1, stroke=0)

        c.setFillColor(colors.black)
        c.setFont('Amiri-Bold' if os.path.exists(FONT_BOLD_FILE) else 'Amiri', 12)
        c.drawString(20*mm, y - 6*mm, ar("بيانات العميل"))

        c.setFont('Amiri', 11)
        client_name = row[2]
        client_phone = ""
        client_address = ""
        client_tax = ""

        if client_info:
            client_name = client_info[0] or row[2]
            client_phone = client_info[1] or ""
            client_address = client_info[3] or ""
            client_tax = client_info[4] or ""

        y -= 14*mm
        c.drawString(20*mm, y, ar(f"الاسم: {client_name}"))
        if client_phone:
            c.drawRightString(width - 20*mm, y, ar(f"الهاتف: {client_phone}"))
        y -= 7*mm
        if client_address:
            c.drawString(20*mm, y, ar(f"العنوان: {client_address}"))
        if client_tax:
            c.drawRightString(width - 20*mm, y, ar(f"الرقم الضريبي: {client_tax}"))
        y -= 7*mm

        y -= 8*mm

        c.setFillColor(colors.HexColor('#0B4F9E'))
        c.rect(15*mm, y - 8*mm, width - 30*mm, 9*mm, fill=1, stroke=0)

        c.setFillColor(colors.white)
        c.setFont('Amiri-Bold' if os.path.exists(FONT_BOLD_FILE) else 'Amiri', 11)
        c.drawCentredString(width - 45*mm, y - 5.5*mm, ar("المبلغ"))
        c.drawCentredString(width / 2, y - 5.5*mm, ar("البيان"))

        y -= 10*mm

        c.setFillColor(colors.black)
        c.setFont('Amiri', 10)
        currency = row[3]

        items = [
            ("إجمالي العهدة المقبوضة", row[4] or 0),
            ("الرسوم الجمركية الأساسية", row[5] or 0),
            ("رسوم الأرضيات والغرامات", row[6] or 0),
            ("مصاريف النقل والتحسين", row[7] or 0),
            ("رسوم الجودة", row[8] or 0),
            ("رسوم الزراعة", row[9] or 0),
            ("رسوم التأمين", row[10] or 0),
            ("مصاريف ورسوم أخرى", row[11] or 0),
        ]

        for cf_name, cf_amount in custom_fees:
            items.append((cf_name, cf_amount or 0))

        items.append(("أتعاب التخليص الجمركي", row[12] or 0))

        alt = False
        for label, amount in items:
            if amount <= 0 and label != "إجمالي العهدة المقبوضة":
                continue

            if alt:
                c.setFillColor(colors.HexColor('#F7F9FC'))
                c.rect(15*mm, y - 4*mm, width - 30*mm, 7*mm, fill=1, stroke=0)
            alt = not alt

            c.setFillColor(colors.black)
            c.setFont('Amiri', 10)
            c.drawString(20*mm, y, ar(label))
            c.setFont('Amiri', 10)
            c.drawRightString(width - 20*mm, y, f"{amount:,.2f} {currency}")
            y -= 7*mm

            if y < 40*mm:
                c.showPage()
                y = height - 30*mm

        y -= 5*mm
        c.setStrokeColor(colors.HexColor('#0B4F9E'))
        c.setLineWidth(1)
        c.line(15*mm, y, width - 15*mm, y)

        y -= 8*mm
        c.setFont('Amiri-Bold' if os.path.exists(FONT_BOLD_FILE) else 'Amiri', 12)
        c.drawString(20*mm, y, ar("إجمالي التكاليف والأتعاب"))
        c.drawRightString(width - 20*mm, y, f"{row[13]:,.2f} {currency}")

        y -= 8*mm
        balance = row[14] or 0

        if balance >= 0:
            c.setFillColor(colors.HexColor('#0B7A3F'))
            c.setFont('Amiri-Bold' if os.path.exists(FONT_BOLD_FILE) else 'Amiri', 13)
            c.drawString(20*mm, y, ar(f"الرصيد المتبقي لكم (دائن): {abs(balance):,.2f} {currency}"))
        else:
            c.setFillColor(colors.HexColor('#C0392B'))
            c.setFont('Amiri-Bold' if os.path.exists(FONT_BOLD_FILE) else 'Amiri', 13)
            c.drawString(20*mm, y, ar(f"المطلوب سداده (مدين): {abs(balance):,.2f} {currency}"))

        c.setFillColor(colors.HexColor('#666666'))
        c.setFont('Amiri', 9)
        c.drawCentredString(width / 2, 20*mm, ar("شكراً لتعاملكم معنا — مكتب أدهم الحاج للتخليص الجمركي"))
        c.drawCentredString(width / 2, 15*mm, ar("هذه فاتورة صادرة إلكترونياً — لا تحتاج إلى توقيع"))

        c.save()

        save_invoice(invoice_no, shipment_no, row[1], client_name,
                     row[13] or 0, currency, filename)

        return filename, invoice_no

    except Exception as e:
        return None, f"خطأ في التوليد: {str(e)}"


# ============================================================
# 6. دوال التصدير والواتساب
# ============================================================
def export_full_excel():
    conn = get_conn()
    cursor = conn.cursor()
    wb = openpyxl.Workbook()

    ws1 = wb.active
    ws1.title = "كشف الشحنات والعهد"
    ws1.append(["رقم الشحنة", "العميل", "العملة", "العهدة المقبوضة", "الرسوم الجمركية",
                "الأرضيات", "النقل", "الجودة", "الزراعة", "التأمين", "أخرى", "الأتعاب",
                "إجمالي التكاليف", "الرصيد المتبقي", "الحالة"])
    cursor.execute("""SELECT shipment_no, client_name, currency, advance_amount, customs_duty,
                      storage_fee, transport_fee, quality_fee, agriculture_fee, insurance_fee,
                      other_fee_amount, clearance_fee, total_expenses, balance, status
                      FROM shipments""")
    for r in cursor.fetchall():
        ws1.append(list(r))

    ws_clients = wb.create_sheet(title="العملاء")
    ws_clients.append(["الاسم", "الهاتف", "البريد", "العنوان", "الرقم الضريبي", "الحد الائتماني"])
    cursor.execute("""SELECT full_name, phone, email, address, tax_number, credit_limit
                      FROM clients""")
    for r in cursor.fetchall():
        ws_clients.append(list(r))

    ws_invoices = wb.create_sheet(title="الفواتير")
    ws_invoices.append(["رقم الفاتورة", "رقم الشحنة", "العميل", "المبلغ", "العملة", "التاريخ"])
    cursor.execute("""SELECT invoice_no, shipment_no, client_name, amount, currency, created_at
                      FROM invoices""")
    for r in cursor.fetchall():
        ws_invoices.append(list(r))

    ws1_5 = wb.create_sheet(title="الرسوم المخصصة")
    ws1_5.append(["رقم الشحنة", "اسم البند", "المبلغ"])
    cursor.execute("SELECT shipment_no, fee_name, amount FROM shipment_custom_fees")
    for r in cursor.fetchall():
        ws1_5.append(list(r))

    ws2 = wb.create_sheet(title="دفتر اليومية العامة ERP")
    ws2.append(["التاريخ", "رقم المرجع", "رمز الحساب", "مدين", "دائن", "العملة", "البيان المحاسبي"])
    cursor.execute(
        "SELECT entry_date, ref_no, account_code, debit, credit, currency, description FROM journal_entries")
    for r in cursor.fetchall():
        ws2.append(list(r))

    ws3 = wb.create_sheet(title="المصروفات")
    ws3.append(["التاريخ", "البيان", "التصنيف", "المبلغ", "العملة"])
    cursor.execute("SELECT exp_date, exp_title, category, amount, currency FROM expenses")
    for r in cursor.fetchall():
        ws3.append(list(r))

    conn.close()
    filename = f"Customs_ERP_Report_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    wb.save(filename)
    return filename


def build_whatsapp_msg(shp_no, include_phone_target=False):
    conn = get_conn()
    cursor = conn.cursor()
    cursor.execute("""SELECT shipment_no, client_name, currency, advance_amount, customs_duty,
                      storage_fee, transport_fee, quality_fee, agriculture_fee, insurance_fee,
                      other_fee_amount, clearance_fee, total_expenses, balance, client_id
                      FROM shipments WHERE shipment_no=?""", (shp_no,))
    row = cursor.fetchone()
    cursor.execute("SELECT fee_name, amount FROM shipment_custom_fees WHERE shipment_no=?", (shp_no,))
    custom_fees = cursor.fetchall()
    conn.close()
    if not row:
        return "الشحنة غير موجودة", None

    bal = row[13]
    state_str = (f"الرصيد المتبقي لكم (دائن): {fmt_num(abs(bal))}" if bal >= 0
                 else f"المبلغ المطلوب سداده منكم (مدين): {fmt_num(abs(bal))}")
    other_str = f"\n- *مصاريف ورسوم أخرى:* {fmt_num(row[10])}" if row[10] > 0 else ""
    custom_str = ""
    for f_name, amt in custom_fees:
        custom_str += f"\n- *{f_name}:* {fmt_num(amt)}"

    msg = f"""*إشعار تسوية حساب شحنة جمركية*
*مكتب أدهم الحاج للتخليص الجمركي والخدمات اللوجستية*
*في جميع المنافذ الجمركية*
-----------------------------------
*المكرم التاجر / العميل:* {row[1]}
*رقم الشحنة:* {row[0]} | *العملة:* {row[2]}
-----------------------------------
*إجمالي العهدة المقبوضة:* {fmt_num(row[3])}
-----------------------------------
*تفاصيل الرسوم والمصاريف الجمركية والميدانية:*
- *الرسوم الجمركية:* {fmt_num(row[4])}
- *الأرضيات والغرامات:* {fmt_num(row[5])}
- *مصاريف النقل والتحسين:* {fmt_num(row[6])}
- *رسوم الجودة:* {fmt_num(row[7])}
- *رسوم الزراعة:* {fmt_num(row[8])}
- *رسوم التأمين:* {fmt_num(row[9])}{other_str}{custom_str}
- *أتعاب التخليص الجمركي:* {fmt_num(row[11])}
-----------------------------------
*إجمالي التكاليف والأتعاب:* {fmt_num(row[12])} {row[2]}
-----------------------------------
*حالة الحساب النهائية:*
*{state_str} {row[2]}*"""

    phone = None
    if include_phone_target and row[14]:
        conn2 = get_conn()
        c2 = conn2.cursor()
        c2.execute("SELECT phone FROM clients WHERE id = ?", (row[14],))
        r2 = c2.fetchone()
        conn2.close()
        if r2 and r2[0]:
            phone = r2[0]

    return msg, phone


def export_client_statement_excel(client_name, date_from=None, date_to=None):
    rows = get_client_statement(client_name, date_from, date_to)
    if not rows:
        return None

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "كشف حساب"

    ws.append([f"كشف حساب العميل: {client_name}"])
    if date_from or date_to:
        ws.append([f"الفترة: من {date_from or 'البداية'} إلى {date_to or 'اليوم'}"])
    ws.append([f"التاريخ: {datetime.now().strftime('%Y-%m-%d %H:%M')}"])
    ws.append([])
    ws.append(["رقم الشحنة", "العملة", "العهد المقبوضة", "إجمالي المصاريف",
               "الأتعاب", "الرصيد", "الحالة", "التاريخ"])

    total_adv = 0.0
    total_exp = 0.0
    total_fee = 0.0
    total_bal = 0.0

    for r in rows:
        ws.append(list(r))
        total_adv += r[2] or 0
        total_exp += r[3] or 0
        total_fee += r[4] or 0
        total_bal += r[5] or 0

    ws.append([])
    ws.append(["الإجمالي", "", total_adv, total_exp, total_fee, total_bal, "", ""])

    filename = f"Statement_{client_name}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    wb.save(filename)
    return filename# ============================================================
# 7. شاشة تسجيل الدخول
# ============================================================
class LoginScreen(Screen):
    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        layout = BoxLayout(orientation='vertical', padding=15, spacing=10)
        lbl_kwargs = {'font_name': ARABIC_FONT} if ARABIC_FONT else {}
        title_kwargs = {'font_name': ARABIC_FONT_BOLD} if ARABIC_FONT_BOLD else {}

        logo_path = get_logo_path()
        if logo_path:
            logo_box = RelativeLayout(size_hint_y=0.42)
            with logo_box.canvas.before:
                Color(1, 1, 1, 1)
                _logo_bg = Rectangle(pos=logo_box.pos, size=logo_box.size)
            logo_box.bind(
                pos=lambda inst, val, bg=_logo_bg: setattr(bg, 'pos', val),
                size=lambda inst, val, bg=_logo_bg: setattr(bg, 'size', val)
            )
            logo_img = Image(
                source=logo_path,
                size_hint=(0.9, 0.9),
                pos_hint={'center_x': 0.5, 'center_y': 0.5},
                allow_stretch=True,
                keep_ratio=True
            )
            logo_box.add_widget(logo_img)
            layout.add_widget(logo_box)

        layout.add_widget(Label(
            text=ar("برنامج التخليص الجمركي اليمني\nتطوير يحيى القداح"),
            font_size='15sp',
            halign='center',
            valign='middle',
            color=(0.2, 0.75, 1, 1),
            size_hint_y=0.10,
            **title_kwargs
        ))

        form_card = BoxLayout(orientation='vertical', padding=10, spacing=12, size_hint_y=0.32)

        row1 = BoxLayout(spacing=10, size_hint_y=None, height=58)
        row1.add_widget(Label(text=ar("اسم المستخدم:"), font_size='15sp', size_hint_x=0.38, **lbl_kwargs))
        self.inp_user = ArabicTextInput(text="admin", hint_text="أدخل اسم المستخدم", size_hint_x=0.62)
        row1.add_widget(self.inp_user)
        form_card.add_widget(row1)

        row2 = BoxLayout(spacing=10, size_hint_y=None, height=58)
        row2.add_widget(Label(text=ar("كلمة المرور:"), font_size='15sp', size_hint_x=0.38, **lbl_kwargs))
        self.inp_pass = ArabicTextInput(text="", password=True, hint_text="أدخل كلمة المرور", size_hint_x=0.62)
        row2.add_widget(self.inp_pass)
        form_card.add_widget(row2)

        layout.add_widget(form_card)

        self.lbl_err = Label(text="", size_hint_y=0.06, color=(1, 0.35, 0.35, 1), **lbl_kwargs)
        layout.add_widget(self.lbl_err)

        btn_login = Button(
            text=ar("تسجيل الدخول إلى النظام المؤسسي"),
            font_size='15sp',
            size_hint_y=0.12,
            background_color=(0.15, 0.45, 0.8, 1),
            background_normal='',
            **title_kwargs
        )
        btn_login.bind(on_press=self.do_login)
        layout.add_widget(btn_login)

        self.add_widget(layout)

    def do_login(self, instance):
        u = self.inp_user.text.strip()
        p = self.inp_pass.text.strip()

        conn = get_conn()
        cursor = conn.cursor()
        cursor.execute("SELECT full_name FROM users WHERE username=? AND password=?", (u, p))
        row = cursor.fetchone()
        conn.close()

        if row:
            self.manager.get_screen('home').set_welcome(row[0])
            self.manager.current = 'home'
            self.lbl_err.text = ""
        else:
            self.lbl_err.text = ar("خطأ: اسم المستخدم أو كلمة المرور غير صحيحة")


# ============================================================
# 8. الشاشة الرئيسية
# ============================================================
class HomeScreen(Screen):
    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.layout = BoxLayout(orientation='vertical', padding=12, spacing=8)
        self.add_widget(self.layout)

    def set_welcome(self, name):
        lbl_kwargs = {'font_name': ARABIC_FONT} if ARABIC_FONT else {}
        title_kwargs = {'font_name': ARABIC_FONT_BOLD} if ARABIC_FONT_BOLD else {}
        self.layout.clear_widgets()

        logo_path = get_logo_path()
        if logo_path:
            logo_box = RelativeLayout(size_hint_y=0.20)
            with logo_box.canvas.before:
                Color(1, 1, 1, 1)
                _bg = Rectangle(pos=logo_box.pos, size=logo_box.size)
            logo_box.bind(
                pos=lambda inst, val, bg=_bg: setattr(bg, 'pos', val),
                size=lambda inst, val, bg=_bg: setattr(bg, 'size', val)
            )
            logo_img = Image(
                source=logo_path,
                size_hint=(0.85, 0.9),
                pos_hint={'center_x': 0.5, 'center_y': 0.5},
                allow_stretch=True,
                keep_ratio=True
            )
            logo_box.add_widget(logo_img)
            self.layout.add_widget(logo_box)

        self.layout.add_widget(Label(
            text=ar(f"مرحباً بك: {name}"),
            font_size='14sp',
            color=(0.3, 0.8, 1, 1),
            size_hint_y=0.045,
            **title_kwargs
        ))

        btns_data = [
            ("تسجيل وتفصيل رسوم الشحنة", "add_shipment", (0.12, 0.45, 0.8, 1)),
            ("سجل الشحنات وتصفية الحسابات", "view_shipments", (0.1, 0.65, 0.4, 1)),
            ("إدارة العملاء", "clients", (0.4, 0.3, 0.65, 1)),
            ("كشف حساب عميل", "statement", (0.55, 0.35, 0.15, 1)),
            ("إدارة الفواتير", "invoices", (0.25, 0.5, 0.6, 1)),
            ("المصروفات التشغيلية والعمومية", "expenses", (0.8, 0.4, 0.15, 1)),
            ("التقرير المالي والقوائم المالية", "dashboard", (0.5, 0.25, 0.7, 1)),
            ("إعدادات الحساب", "settings", (0.3, 0.3, 0.5, 1)),
        ]

        for text, scr_name, col in btns_data:
            btn = Button(text=ar(text), font_size='12sp', size_hint_y=0.063,
                         background_color=col, background_normal='', **lbl_kwargs)
            btn.bind(on_press=lambda x, s=scr_name: self.navigate(s))
            self.layout.add_widget(btn)

        btn_backup = Button(
            text=ar("نسخة احتياطية"),
            font_size='12sp',
            size_hint_y=0.06,
            background_color=(0.2, 0.55, 0.55, 1),
            background_normal='',
            **lbl_kwargs
        )
        btn_backup.bind(on_press=self.do_backup)
        self.layout.add_widget(btn_backup)

        b_box = BoxLayout(size_hint_y=0.06, spacing=5)
        btn_logout = Button(text=ar("تسجيل الخروج"),
                            background_color=(0.8, 0.2, 0.2, 1),
                            background_normal='', **lbl_kwargs)
        btn_logout.bind(on_press=lambda x: setattr(self.manager, 'current', 'login'))
        b_box.add_widget(btn_logout)
        self.layout.add_widget(b_box)

    def do_backup(self, instance):
        result = backup_database()
        lbl_kwargs = {'font_name': ARABIC_FONT} if ARABIC_FONT else {}
        Popup(
            title=ar("النسخ الاحتياطي"),
            content=Label(text=ar(result), **lbl_kwargs),
            size_hint=(0.85, 0.25)
        ).open()

    def navigate(self, screen_name):
        if screen_name == 'view_shipments':
            self.manager.get_screen('view_shipments').load_data()
        elif screen_name == 'clients':
            self.manager.get_screen('clients').load_clients()
        elif screen_name == 'invoices':
            self.manager.get_screen('invoices').load_invoices()
        elif screen_name == 'expenses':
            self.manager.get_screen('expenses').load_categories()
        elif screen_name == 'dashboard':
            self.manager.get_screen('dashboard').load_main_menu()
        elif screen_name == 'settings':
            self.manager.get_screen('settings').load_settings()
        elif screen_name == 'add_shipment':
            self.manager.get_screen('add_shipment').reset_form_mode()
        elif screen_name == 'statement':
            self.manager.get_screen('statement').load_clients()
        self.manager.current = screen_name


# ============================================================
# 9. شاشة الإعدادات
# ============================================================
class SettingsScreen(Screen):
    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.layout = BoxLayout(orientation='vertical', padding=15, spacing=10)
        self.add_widget(self.layout)

    def load_settings(self):
        lbl_kwargs = {'font_name': ARABIC_FONT} if ARABIC_FONT else {}
        title_kwargs = {'font_name': ARABIC_FONT_BOLD} if ARABIC_FONT_BOLD else {}
        self.layout.clear_widgets()

        logo_path = get_logo_path()
        if logo_path:
            logo_box = RelativeLayout(size_hint_y=0.22)
            with logo_box.canvas.before:
                Color(1, 1, 1, 1)
                _bg = Rectangle(pos=logo_box.pos, size=logo_box.size)
            logo_box.bind(
                pos=lambda inst, val, bg=_bg: setattr(bg, 'pos', val),
                size=lambda inst, val, bg=_bg: setattr(bg, 'size', val)
            )
            logo_box.add_widget(Image(
                source=logo_path,
                size_hint=(0.85, 0.9),
                pos_hint={'center_x': 0.5, 'center_y': 0.5},
                allow_stretch=True,
                keep_ratio=True
            ))
            self.layout.add_widget(logo_box)

        self.layout.add_widget(Label(
            text=ar("إعدادات الحساب وتغيير بيانات الدخول"),
            font_size='16sp', color=(0.2, 0.75, 1, 1),
            size_hint_y=0.07, **title_kwargs))

        form = BoxLayout(orientation='vertical', spacing=12, size_hint_y=0.5)

        row_old = BoxLayout(spacing=10, size_hint_y=None, height=60)
        row_old.add_widget(Label(text=ar("كلمة المرور الحالية:"), font_size='14sp',
                                 size_hint_x=0.4, **lbl_kwargs))
        self.inp_old_pass = ArabicTextInput(password=True, hint_text="أدخل كلمة المرور الحالية", size_hint_x=0.6)
        row_old.add_widget(self.inp_old_pass)
        form.add_widget(row_old)

        row1 = BoxLayout(spacing=10, size_hint_y=None, height=60)
        row1.add_widget(Label(text=ar("اسم المستخدم الجديد:"), font_size='14sp',
                              size_hint_x=0.4, **lbl_kwargs))
        self.inp_new_user = ArabicTextInput(hint_text="أدخل اسم المستخدم الجديد", size_hint_x=0.6)
        row1.add_widget(self.inp_new_user)
        form.add_widget(row1)

        row2 = BoxLayout(spacing=10, size_hint_y=None, height=60)
        row2.add_widget(Label(text=ar("كلمة المرور الجديدة:"), font_size='14sp',
                              size_hint_x=0.4, **lbl_kwargs))
        self.inp_new_pass = ArabicTextInput(password=True, hint_text="أدخل كلمة المرور الجديدة", size_hint_x=0.6)
        row2.add_widget(self.inp_new_pass)
        form.add_widget(row2)

        self.layout.add_widget(form)

        self.lbl_status = Label(text="", size_hint_y=0.08, **lbl_kwargs)
        self.layout.add_widget(self.lbl_status)

        btns = BoxLayout(size_hint_y=0.13, spacing=10)
        btn_save = Button(text=ar("حفظ بيانات الدخول الجديدة"),
                          background_color=(0, 0.6, 0.3, 1),
                          background_normal='', **lbl_kwargs)
        btn_save.bind(on_press=self.save_credentials)
        btn_back = Button(text=ar("الرئيسية"),
                          background_color=(0.35, 0.4, 0.5, 1),
                          background_normal='', **lbl_kwargs)
        btn_back.bind(on_press=lambda x: setattr(self.manager, 'current', 'home'))

        btns.add_widget(btn_save)
        btns.add_widget(btn_back)
        self.layout.add_widget(btns)

    def save_credentials(self, instance):
        old_p = self.inp_old_pass.text.strip()
        new_u = self.inp_new_user.text.strip()
        new_p = self.inp_new_pass.text.strip()

        if not old_p or not new_u or not new_p:
            self.lbl_status.text = ar("يرجى تعبئة جميع الحقول")
            self.lbl_status.color = (1, 0, 0, 1)
            return

        conn = get_conn()
        cursor = conn.cursor()
        cursor.execute("SELECT password FROM users LIMIT 1")
        row = cursor.fetchone()

        if not row or row[0] != old_p:
            self.lbl_status.text = ar("خطأ: كلمة المرور الحالية غير صحيحة!")
            self.lbl_status.color = (1, 0, 0, 1)
            conn.close()
            return

        try:
            cursor.execute("UPDATE users SET username=?, password=? WHERE password=?",
                           (new_u, new_p, old_p))
            if cursor.rowcount == 0:
                cursor.execute("INSERT OR REPLACE INTO users (username, password, full_name) VALUES (?, ?, ?)",
                               (new_u, new_p, "أدهم الحاج"))
            conn.commit()
            self.lbl_status.text = ar("تم تحديث بيانات الدخول بنجاح!")
            self.lbl_status.color = (0, 1, 0, 1)
        except Exception as e:
            self.lbl_status.text = ar(f"خطأ: {str(e)}")
            self.lbl_status.color = (1, 0, 0, 1)
        finally:
            conn.close()


# ============================================================
# 10. بطاقة العميل
# ============================================================
class ClientCard(BoxLayout):
    def __init__(self, client_data, on_edit, on_delete, on_whatsapp, on_statement, **kwargs):
        super().__init__(orientation='vertical', size_hint_y=None, height=170,
                         spacing=5, padding=(12, 8, 12, 8), **kwargs)

        lbl_kwargs = {'font_name': ARABIC_FONT} if ARABIC_FONT else {}
        title_kwargs = {'font_name': ARABIC_FONT_BOLD} if ARABIC_FONT_BOLD else {}

        with self.canvas.before:
            Color(0.09, 0.12, 0.17, 1)
            _bg = RoundedRectangle(pos=self.pos, size=self.size, radius=[10])
        self.bind(
            pos=lambda inst, val, bg=_bg: setattr(bg, 'pos', val),
            size=lambda inst, val, bg=_bg: setattr(bg, 'size', val))

        cid, name, phone, email, address, tax_num, credit_limit, notes, created = client_data

        header = BoxLayout(size_hint_y=None, height=32, spacing=8)

        lbl_name = Label(
            text=ar(f"{name}"),
            font_size='15sp',
            color=(1, 1, 1, 1),
            halign='left', valign='middle',
            size_hint_x=0.55,
            **title_kwargs)
        lbl_name.bind(size=lambda inst, val: setattr(inst, 'text_size', val))

        phone_str = phone if phone else "بدون هاتف"
        lbl_phone = Label(
            text=ar(f"{phone_str}"),
            font_size='12sp',
            color=(0.4, 0.85, 0.5, 1) if phone else (0.6, 0.5, 0.5, 1),
            halign='right', valign='middle',
            size_hint_x=0.45,
            **lbl_kwargs)
        lbl_phone.bind(size=lambda inst, val: setattr(inst, 'text_size', val))

        header.add_widget(lbl_name)
        header.add_widget(lbl_phone)
        self.add_widget(header)

        sep = BoxLayout(size_hint_y=None, height=1)
        with sep.canvas.before:
            Color(0.2, 0.28, 0.38, 1)
            _sep = Rectangle(pos=sep.pos, size=sep.size)
        sep.bind(
            pos=lambda inst, val, s=_sep: setattr(s, 'pos', val),
            size=lambda inst, val, s=_sep: setattr(s, 'size', val))
        self.add_widget(sep)

        if address:
            lbl_addr = Label(
                text=ar(f"{address}"),
                font_size='11sp',
                color=(0.7, 0.75, 0.85, 1),
                halign='right', valign='middle',
                size_hint_y=None, height=20,
                **lbl_kwargs)
            lbl_addr.bind(size=lambda inst, val: setattr(inst, 'text_size', (val[0] - 4, val[1])))
            self.add_widget(lbl_addr)

        stats = get_client_stats(cid)
        stats_text = f"الشحنات: {stats['count']}  |  الرصيد: {fmt_num(stats['balance'])}"
        lbl_stats = Label(
            text=ar(stats_text),
            font_size='11sp',
            color=(1, 0.75, 0.3, 1),
            halign='right', valign='middle',
            size_hint_y=None, height=20,
            **lbl_kwargs)
        lbl_stats.bind(size=lambda inst, val: setattr(inst, 'text_size', (val[0] - 4, val[1])))
        self.add_widget(lbl_stats)

        bottom = BoxLayout(size_hint_y=None, height=40, spacing=5)

        btn_wa = Button(
            text=ar("واتساب"),
            font_size='11sp',
            background_color=(0, 0.65, 0.35, 1) if phone else (0.4, 0.4, 0.4, 1),
            background_normal='',
            size_hint_x=0.22,
            disabled=not phone,
            **lbl_kwargs)
        btn_wa.bind(on_press=lambda x: on_whatsapp(phone, name))

        btn_stmt = Button(
            text=ar("كشف"),
            font_size='11sp',
            background_color=(0.55, 0.35, 0.15, 1),
            background_normal='',
            size_hint_x=0.18,
            **lbl_kwargs)
        btn_stmt.bind(on_press=lambda x: on_statement(name))

        btn_edit = Button(
            text=ar("تعديل"),
            font_size='11sp',
            background_color=(0.15, 0.45, 0.8, 1),
            background_normal='',
            size_hint_x=0.18,
            **lbl_kwargs)
        btn_edit.bind(on_press=lambda x: on_edit(cid))

        btn_del = Button(
            text=ar("حذف"),
            font_size='11sp',
            background_color=(0.75, 0.2, 0.2, 1),
            background_normal='',
            size_hint_x=0.18,
            **lbl_kwargs)
        btn_del.bind(on_press=lambda x: on_delete(cid))

        lbl_count = Label(
            text=ar(f"({stats['count']})"),
            font_size='11sp',
            color=(0.5, 0.6, 0.75, 1),
            halign='right', valign='middle',
            size_hint_x=0.24,
            **lbl_kwargs)
        lbl_count.bind(size=lambda inst, val: setattr(inst, 'text_size', val))

        bottom.add_widget(btn_del)
        bottom.add_widget(btn_edit)
        bottom.add_widget(btn_stmt)
        bottom.add_widget(btn_wa)
        bottom.add_widget(lbl_count)
        self.add_widget(bottom)


# ============================================================
# 11. شاشة إدارة العملاء
# ============================================================
class ClientsScreen(Screen):
    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.layout = BoxLayout(orientation='vertical', padding=(10, 12, 10, 8), spacing=6)
        self.add_widget(self.layout)

    def load_clients(self):
        lbl_kwargs = {'font_name': ARABIC_FONT} if ARABIC_FONT else {}
        title_kwargs = {'font_name': ARABIC_FONT_BOLD} if ARABIC_FONT_BOLD else {}
        self.layout.clear_widgets()

        title_box = BoxLayout(size_hint_y=None, height=40, spacing=6)
        title_box.add_widget(Label(
            text=ar("إدارة العملاء"),
            font_size='17sp',
            color=(0.2, 0.8, 1, 1),
            **title_kwargs))

        btn_add = Button(
            text=ar("+ عميل جديد"),
            background_color=(0.15, 0.6, 0.4, 1),
            background_normal='',
            font_size='13sp',
            size_hint_x=0.35,
            **lbl_kwargs)
        btn_add.bind(on_press=lambda x: self.open_client_form(None))
        title_box.add_widget(btn_add)
        self.layout.add_widget(title_box)

        search_box = BoxLayout(size_hint_y=None, height=42, spacing=6)
        search_box.add_widget(Label(
            text=ar("بحث:"),
            font_size='13sp',
            size_hint_x=0.15,
            halign='right',
            valign='middle',
            **lbl_kwargs))

        self.inp_search = ArabicTextInput(
            hint_text="اكتب اسم العميل أو رقم الهاتف...",
            size_hint_x=0.85,
            size_hint_y=None,
            height=42)
        self.inp_search.real_input.bind(text=lambda s, t: self.filter_clients(t))
        search_box.add_widget(self.inp_search)
        self.layout.add_widget(search_box)

        scroll = ScrollView(size_hint=(1, 1))
        self.list_box = BoxLayout(
            orientation='vertical',
            size_hint_y=None,
            spacing=10,
            padding=(2, 4))
        self.list_box.bind(minimum_height=self.list_box.setter('height'))

        clients = get_all_clients_objects()

        if not clients:
            empty_box = BoxLayout(orientation='vertical', size_hint_y=None, height=100)
            empty_box.add_widget(Label(
                text=ar("لا يوجد عملاء مسجلون"),
                font_size='14sp',
                color=(0.6, 0.65, 0.75, 1),
                **lbl_kwargs))
            self.list_box.add_widget(empty_box)
        else:
            for c in clients:
                card = ClientCard(c, self.open_client_form, self.delete_client,
                                  self.send_whatsapp_to_phone, self.open_statement)
                self.list_box.add_widget(card)

        scroll.add_widget(self.list_box)
        self.layout.add_widget(scroll)

        btns = BoxLayout(size_hint_y=None, height=45, spacing=6)
        btn_back = Button(
            text=ar("الرئيسية"),
            background_color=(0.35, 0.4, 0.5, 1),
            background_normal='',
            font_size='13sp',
            **lbl_kwargs)
        btn_back.bind(on_press=lambda x: setattr(self.manager, 'current', 'home'))
        btns.add_widget(btn_back)
        self.layout.add_widget(btns)

    def filter_clients(self, text):
        if not hasattr(self, 'list_box'):
            return
        lbl_kwargs = {'font_name': ARABIC_FONT} if ARABIC_FONT else {}
        self.list_box.clear_widgets()
        clients = get_all_clients_objects()
        if text:
            t = text.strip().lower()
            clients = [c for c in clients
                       if t in c[1].lower() or (c[2] and t in c[2].lower())]

        if not clients:
            self.list_box.add_widget(Label(
                text=ar("لا توجد نتائج"),
                font_size='14sp',
                color=(0.6, 0.65, 0.75, 1),
                size_hint_y=None, height=60,
                **lbl_kwargs))
            return

        for c in clients:
            card = ClientCard(c, self.open_client_form, self.delete_client,
                              self.send_whatsapp_to_phone, self.open_statement)
            self.list_box.add_widget(card)

    def open_client_form(self, client_id):
        lbl_kwargs = {'font_name': ARABIC_FONT} if ARABIC_FONT else {}
        title_kwargs = {'font_name': ARABIC_FONT_BOLD} if ARABIC_FONT_BOLD else {}

        client = None
        if client_id:
            client = get_client_by_id(client_id)

        box = BoxLayout(orientation='vertical', padding=12, spacing=6)
        box.add_widget(Label(
            text=ar("تعديل بيانات العميل") if client else ar("إضافة عميل جديد"),
            font_size='15sp',
            size_hint_y=None, height=30,
            color=(0.3, 0.85, 1, 1),
            **title_kwargs))

        inp_name = ArabicTextInput(
            text=client[1] if client else '',
            hint_text="اسم العميل الكامل *",
            size_hint_y=None, height=48)
        box.add_widget(inp_name)

        inp_phone = ArabicTextInput(
            text=client[2] if client and client[2] else '',
            hint_text="رقم الهاتف (مثال: 967712345678)",
            size_hint_y=None, height=48)
        box.add_widget(inp_phone)

        inp_email = ArabicTextInput(
            text=client[3] if client and client[3] else '',
            hint_text="البريد الإلكتروني",
            size_hint_y=None, height=48)
        box.add_widget(inp_email)

        inp_address = ArabicTextInput(
            text=client[4] if client and client[4] else '',
            hint_text="العنوان",
            size_hint_y=None, height=48)
        box.add_widget(inp_address)

        inp_tax = ArabicTextInput(
            text=client[5] if client and client[5] else '',
            hint_text="الرقم الضريبي",
            size_hint_y=None, height=48)
        box.add_widget(inp_tax)

        inp_credit = ArabicTextInput(
            text=str(client[6]) if client and client[6] else '0',
            hint_text="الحد الائتماني",
            input_filter='float',
            size_hint_y=None, height=48)
        box.add_widget(inp_credit)

        inp_notes = ArabicTextInput(
            text=client[7] if client and client[7] else '',
            hint_text="ملاحظات",
            size_hint_y=None, height=48)
        box.add_widget(inp_notes)

        btn_row = BoxLayout(size_hint_y=None, height=45, spacing=8)
        btn_save = Button(
            text=ar("حفظ"),
            background_color=(0, 0.6, 0.3, 1),
            background_normal='',
            **lbl_kwargs)
        btn_cancel = Button(
            text=ar("إلغاء"),
            background_color=(0.4, 0.4, 0.5, 1),
            background_normal='',
            **lbl_kwargs)

        popup = Popup(
            title=ar("بيانات العميل"),
            content=box,
            size_hint=(0.92, 0.78),
            title_font=ARABIC_FONT_BOLD if ARABIC_FONT_BOLD else 'Roboto',
        )

        def do_save(inst):
            name = inp_name.text.strip()
            if not name:
                return
            phone = inp_phone.text.strip()
            email = inp_email.text.strip()
            address = inp_address.text.strip()
            tax = inp_tax.text.strip()
            try:
                credit = float(inp_credit.text or 0)
            except ValueError:
                credit = 0.0
            notes = inp_notes.text.strip()

            ok, msg = save_client(client_id, name, phone, email, address,
                                   tax, credit, notes)
            if ok:
                popup.dismiss()
                self.load_clients()
            else:
                lbl_kwargs2 = {'font_name': ARABIC_FONT} if ARABIC_FONT else {}
                Popup(
                    title=ar("خطأ"),
                    content=Label(text=ar(msg), **lbl_kwargs2),
                    size_hint=(0.85, 0.22)
                ).open()

        btn_save.bind(on_press=do_save)
        btn_cancel.bind(on_press=lambda x: popup.dismiss())

        btn_row.add_widget(btn_cancel)
        btn_row.add_widget(btn_save)
        box.add_widget(btn_row)

        popup.open()

    def delete_client(self, client_id):
        lbl_kwargs = {'font_name': ARABIC_FONT} if ARABIC_FONT else {}
        title_kwargs = {'font_name': ARABIC_FONT_BOLD} if ARABIC_FONT_BOLD else {}

        client = get_client_by_id(client_id)
        if not client:
            return

        box = BoxLayout(orientation='vertical', padding=15, spacing=10)
        box.add_widget(Label(
            text=ar(f"حذف العميل: {client[1]}؟"),
            font_size='15sp',
            halign='center',
            size_hint_y=None, height=30,
            **title_kwargs))

        box.add_widget(Label(
            text=ar("لا يمكن الحذف إذا كان العميل لديه شحنات"),
            font_size='12sp',
            halign='center',
            color=(0.9, 0.6, 0.3, 1),
            size_hint_y=None, height=26,
            **lbl_kwargs))

        btn_row = BoxLayout(size_hint_y=None, height=45, spacing=8)
        btn_yes = Button(
            text=ar("حذف"),
            background_color=(0.7, 0.2, 0.2, 1),
            background_normal='',
            **lbl_kwargs)
        btn_no = Button(
            text=ar("إلغاء"),
            background_color=(0.3, 0.35, 0.45, 1),
            background_normal='',
            **lbl_kwargs)

        popup = Popup(
            title=ar("تأكيد الحذف"),
            content=box,
            size_hint=(0.85, 0.35),
            title_font=ARABIC_FONT_BOLD if ARABIC_FONT_BOLD else 'Roboto',
        )

        def confirm_del(inst):
            ok, msg = delete_client(client_id)
            popup.dismiss()
            self.load_clients()
            if not ok:
                lbl_kwargs2 = {'font_name': ARABIC_FONT} if ARABIC_FONT else {}
                Popup(
                    title=ar("تعذر الحذف"),
                    content=Label(text=ar(msg), **lbl_kwargs2),
                    size_hint=(0.85, 0.25)
                ).open()

        btn_yes.bind(on_press=confirm_del)
        btn_no.bind(on_press=lambda x: popup.dismiss())

        btn_row.add_widget(btn_yes)
        btn_row.add_widget(btn_no)
        box.add_widget(btn_row)
        popup.open()

    def send_whatsapp_to_phone(self, phone, client_name):
        if not phone:
            return
        msg = f"السلام عليكم {client_name},\nهذه رسالة من مكتب أدهم الحاج للتخليص الجمركي."
        encoded = urllib.parse.quote(msg)
        clean_phone = ''.join(c for c in phone if c.isdigit())
        url = f"whatsapp://send?phone={clean_phone}&text={encoded}"
        try:
            webbrowser.open(url)
        except Exception:
            try:
                webbrowser.open(f"https://wa.me/{clean_phone}?text={encoded}")
            except Exception:
                pass

    def open_statement(self, client_name):
        statement_screen = self.manager.get_screen('statement')
        statement_screen.load_clients()
        try:
            target_ar = ar(client_name)
            if target_ar in statement_screen.spn_client.values:
                statement_screen.spn_client.text = target_ar
        except Exception:
            pass
        self.manager.current = 'statement'# ============================================================
# 12. بطاقة الشحنة
# ============================================================
class ShipmentCard(BoxLayout):
    def __init__(self, shp_data, on_edit, on_delete, on_whatsapp, on_statement,
                 on_invoice, **kwargs):
        super().__init__(orientation='vertical', size_hint_y=None, height=200,
                         spacing=6, padding=(12, 10, 12, 10), **kwargs)

        lbl_kwargs = {'font_name': ARABIC_FONT} if ARABIC_FONT else {}
        title_kwargs = {'font_name': ARABIC_FONT_BOLD} if ARABIC_FONT_BOLD else {}

        with self.canvas.before:
            Color(0.09, 0.12, 0.17, 1)
            _bg = RoundedRectangle(pos=self.pos, size=self.size, radius=[10])
        self.bind(
            pos=lambda inst, val, bg=_bg: setattr(bg, 'pos', val),
            size=lambda inst, val, bg=_bg: setattr(bg, 'size', val)
        )

        shp_no, client, curr, adv, total_exp, fee, balance, status = shp_data

        header = BoxLayout(size_hint_y=None, height=36, spacing=8)

        lbl_no = Label(
            text=ar(f"شحنة رقم: {shp_no}"),
            font_size='15sp',
            color=(0.3, 0.85, 1, 1),
            halign='left', valign='middle',
            size_hint_x=0.5,
            **title_kwargs
        )
        lbl_no.bind(size=lambda inst, val: setattr(inst, 'text_size', val))

        lbl_client = Label(
            text=ar(f"العميل: {client}"),
            font_size='14sp',
            color=(1, 1, 1, 1),
            halign='right', valign='middle',
            size_hint_x=0.5,
            **lbl_kwargs
        )
        lbl_client.bind(size=lambda inst, val: setattr(inst, 'text_size', val))

        header.add_widget(lbl_no)
        header.add_widget(lbl_client)
        self.add_widget(header)

        sep = BoxLayout(size_hint_y=None, height=1)
        with sep.canvas.before:
            Color(0.25, 0.32, 0.42, 1)
            _sep = Rectangle(pos=sep.pos, size=sep.size)
        sep.bind(
            pos=lambda inst, val, s=_sep: setattr(s, 'pos', val),
            size=lambda inst, val, s=_sep: setattr(s, 'size', val)
        )
        self.add_widget(sep)

        row1 = BoxLayout(size_hint_y=None, height=26, spacing=8)

        r1_left = Label(
            text=ar(f"العهد المقبوضة:  {fmt_num(adv)} {curr}"),
            font_size='12sp',
            color=(0.4, 0.7, 1, 1),
            halign='left', valign='middle',
            size_hint_x=0.5,
            **lbl_kwargs
        )
        r1_left.bind(size=lambda inst, val: setattr(inst, 'text_size', val))

        r1_right = Label(
            text=ar(f"إجمالي المصاريف:  {fmt_num(total_exp)} {curr}"),
            font_size='12sp',
            color=(1, 0.7, 0.3, 1),
            halign='right', valign='middle',
            size_hint_x=0.5,
            **lbl_kwargs
        )
        r1_right.bind(size=lambda inst, val: setattr(inst, 'text_size', val))

        row1.add_widget(r1_left)
        row1.add_widget(r1_right)
        self.add_widget(row1)

        if balance >= 0:
            bal_color = (0.3, 0.9, 0.5, 1)
            bal_text = f"الرصيد المتبقي (دائن):  {fmt_num(abs(balance))} {curr}"
        else:
            bal_color = (1, 0.4, 0.4, 1)
            bal_text = f"المطلوب سداده (مدين):  {fmt_num(abs(balance))} {curr}"

        lbl_bal = Label(
            text=ar(bal_text),
            font_size='13sp',
            color=bal_color,
            halign='right', valign='middle',
            size_hint_y=None, height=28,
            **title_kwargs
        )
        lbl_bal.bind(size=lambda inst, val: setattr(inst, 'text_size', (val[0] - 4, val[1])))
        self.add_widget(lbl_bal)

        bottom = BoxLayout(size_hint_y=None, height=42, spacing=4)

        btn_inv = Button(
            text=ar("فاتورة"),
            font_size='10sp',
            background_color=(0.7, 0.5, 0.15, 1),
            background_normal='',
            size_hint_x=0.18,
            **lbl_kwargs
        )
        btn_inv.bind(on_press=lambda x: on_invoice(shp_no))

        btn_stmt = Button(
            text=ar("كشف"),
            font_size='10sp',
            background_color=(0.55, 0.35, 0.15, 1),
            background_normal='',
            size_hint_x=0.13,
            **lbl_kwargs
        )
        btn_stmt.bind(on_press=lambda x: on_statement(client))

        btn_wa = Button(
            text=ar("واتساب"),
            font_size='10sp',
            background_color=(0, 0.65, 0.35, 1),
            background_normal='',
            size_hint_x=0.18,
            **lbl_kwargs
        )
        btn_wa.bind(on_press=lambda x: on_whatsapp(shp_no))

        btn_del = Button(
            text=ar("حذف"),
            font_size='10sp',
            background_color=(0.75, 0.2, 0.2, 1),
            background_normal='',
            size_hint_x=0.13,
            **lbl_kwargs
        )
        btn_del.bind(on_press=lambda x: on_delete(shp_no))

        btn_edit = Button(
            text=ar("تعديل"),
            font_size='10sp',
            background_color=(0.15, 0.45, 0.8, 1),
            background_normal='',
            size_hint_x=0.13,
            **lbl_kwargs
        )
        btn_edit.bind(on_press=lambda x: on_edit(shp_no))

        status_clean = "قيد الإجراء" if "تحت" in status else "مكتملة ومسلمة"
        if "مكتملة" in status:
            status_color = (0.3, 0.9, 0.5, 1)
        else:
            status_color = (1, 0.75, 0.3, 1)

        lbl_status = Label(
            text=ar(f"الحالة: {status_clean}"),
            font_size='10sp',
            color=status_color,
            halign='right', valign='middle',
            size_hint_x=0.25,
            **lbl_kwargs
        )
        lbl_status.bind(size=lambda inst, val: setattr(inst, 'text_size', val))

        bottom.add_widget(btn_inv)
        bottom.add_widget(btn_stmt)
        bottom.add_widget(btn_wa)
        bottom.add_widget(btn_del)
        bottom.add_widget(btn_edit)
        bottom.add_widget(lbl_status)
        self.add_widget(bottom)


# ============================================================
# 13. شاشة سجل الشحنات
# ============================================================
class ViewShipmentsScreen(Screen):
    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.layout = BoxLayout(orientation='vertical', padding=(10, 12, 10, 8), spacing=6)
        self.add_widget(self.layout)

    def load_data(self):
        lbl_kwargs = {'font_name': ARABIC_FONT} if ARABIC_FONT else {}
        title_kwargs = {'font_name': ARABIC_FONT_BOLD} if ARABIC_FONT_BOLD else {}
        self.layout.clear_widgets()

        title_box = BoxLayout(size_hint_y=None, height=40)
        title_box.add_widget(Label(
            text=ar("سجل الشحنات"),
            font_size='17sp',
            color=(0.2, 0.8, 1, 1),
            **title_kwargs))
        self.layout.add_widget(title_box)

        scroll = ScrollView(size_hint=(1, 1))
        self.list_box = BoxLayout(
            orientation='vertical',
            size_hint_y=None,
            spacing=10,
            padding=(2, 4)
        )
        self.list_box.bind(minimum_height=self.list_box.setter('height'))

        conn = get_conn()
        cursor = conn.cursor()
        cursor.execute("""SELECT shipment_no, client_name, currency, advance_amount,
                          total_expenses, clearance_fee, balance, status FROM shipments
                          ORDER BY id DESC""")
        rows = cursor.fetchall()
        conn.close()

        if not rows:
            empty_box = BoxLayout(orientation='vertical', size_hint_y=None, height=100)
            empty_box.add_widget(Label(
                text=ar("لا توجد شحنات مسجلة حالياً"),
                font_size='14sp',
                color=(0.6, 0.65, 0.75, 1),
                **lbl_kwargs))
            self.list_box.add_widget(empty_box)
        else:
            for r in rows:
                card = ShipmentCard(r, self.edit_shipment, self.delete_shipment,
                                    self.send_whatsapp, self.open_statement,
                                    self.generate_invoice)
                self.list_box.add_widget(card)

        scroll.add_widget(self.list_box)
        self.layout.add_widget(scroll)

        btns = BoxLayout(size_hint_y=None, height=45, spacing=6)

        btn_excel = Button(
            text=ar("Excel"),
            background_color=(0.1, 0.45, 0.3, 1),
            background_normal='',
            font_size='13sp',
            **lbl_kwargs)
        btn_excel.bind(on_press=self.export_excel)

        btn_back = Button(
            text=ar("الرئيسية"),
            background_color=(0.35, 0.4, 0.5, 1),
            background_normal='',
            font_size='13sp',
            **lbl_kwargs)
        btn_back.bind(on_press=lambda x: setattr(self.manager, 'current', 'home'))

        btns.add_widget(btn_back)
        btns.add_widget(btn_excel)
        self.layout.add_widget(btns)

    def export_excel(self, instance):
        try:
            export_full_excel()
        except Exception:
            pass

    def edit_shipment(self, shp_no):
        add_screen = self.manager.get_screen('add_shipment')
        add_screen.reset_form_mode()
        add_screen.load_shipment_data_for_edit(shp_no)
        self.manager.current = 'add_shipment'

    def send_whatsapp(self, shp_no):
        msg, phone = build_whatsapp_msg(shp_no, include_phone_target=True)
        if msg == "الشحنة غير موجودة":
            return
        encoded_msg = urllib.parse.quote(msg)

        if phone:
            clean_phone = ''.join(c for c in phone if c.isdigit())
            url = f"whatsapp://send?phone={clean_phone}&text={encoded_msg}"
            try:
                webbrowser.open(url)
                return
            except Exception:
                try:
                    webbrowser.open(f"https://wa.me/{clean_phone}?text={encoded_msg}")
                    return
                except Exception:
                    pass

        try:
            webbrowser.open(f"whatsapp://send?text={encoded_msg}")
        except Exception:
            pass

    def open_statement(self, client_name):
        statement_screen = self.manager.get_screen('statement')
        statement_screen.load_clients()
        try:
            target_ar = ar(client_name)
            if target_ar in statement_screen.spn_client.values:
                statement_screen.spn_client.text = target_ar
        except Exception:
            pass
        self.manager.current = 'statement'

    def generate_invoice(self, shp_no):
        lbl_kwargs = {'font_name': ARABIC_FONT} if ARABIC_FONT else {}
        title_kwargs = {'font_name': ARABIC_FONT_BOLD} if ARABIC_FONT_BOLD else {}

        if not PDF_AVAILABLE:
            Popup(
                title=ar("غير متاح"),
                content=Label(
                    text=ar("مكتبة reportlab غير مثبتة."),
                    halign='center',
                    **lbl_kwargs),
                size_hint=(0.9, 0.3)
            ).open()
            return

        box = BoxLayout(orientation='vertical', padding=15, spacing=10)
        box.add_widget(Label(
            text=ar("جاري توليد الفاتورة..."),
            font_size='14sp',
            halign='center',
            size_hint_y=None, height=40,
            **title_kwargs))

        popup = Popup(
            title=ar("الفاتورة"),
            content=box,
            size_hint=(0.85, 0.25),
            title_font=ARABIC_FONT_BOLD if ARABIC_FONT_BOLD else 'Roboto',
            auto_dismiss=False
        )
        popup.open()

        def do_generate(dt):
            popup.dismiss()
            pdf_path, invoice_no = generate_invoice_pdf(shp_no)

            if not pdf_path:
                Popup(
                    title=ar("خطأ"),
                    content=Label(
                        text=ar(f"فشل توليد الفاتورة:\n{invoice_no}"),
                        halign='center',
                        **lbl_kwargs),
                    size_hint=(0.85, 0.3)
                ).open()
                return

            box2 = BoxLayout(orientation='vertical', padding=15, spacing=10)

            box2.add_widget(Label(
                text=ar(f"تم إنشاء الفاتورة: {invoice_no}"),
                font_size='14sp',
                halign='center',
                color=(0.3, 0.9, 0.5, 1),
                size_hint_y=None, height=30,
                **title_kwargs))

            box2.add_widget(Label(
                text=ar(f"المسار: {pdf_path}"),
                font_size='11sp',
                halign='center',
                color=(0.7, 0.75, 0.85, 1),
                size_hint_y=None, height=26,
                **lbl_kwargs))

            btn_row = BoxLayout(size_hint_y=None, height=45, spacing=6)

            btn_open = Button(
                text=ar("فتح PDF"),
                background_color=(0.15, 0.45, 0.8, 1),
                background_normal='',
                **lbl_kwargs)

            btn_share = Button(
                text=ar("مشاركة واتساب"),
                background_color=(0, 0.65, 0.35, 1),
                background_normal='',
                **lbl_kwargs)

            btn_close = Button(
                text=ar("إغلاق"),
                background_color=(0.4, 0.4, 0.5, 1),
                background_normal='',
                **lbl_kwargs)

            popup2 = Popup(
                title=ar("نجاح"),
                content=box2,
                size_hint=(0.9, 0.4),
                title_font=ARABIC_FONT_BOLD if ARABIC_FONT_BOLD else 'Roboto',
            )

            def open_pdf(inst):
                try:
                    if os.name == 'nt':
                        os.startfile(pdf_path)
                    else:
                        webbrowser.open(f"file://{os.path.abspath(pdf_path)}")
                except Exception:
                    pass

            def share_wa(inst):
                msg, phone = build_whatsapp_msg(shp_no, include_phone_target=True)
                share_msg = f"مرفق فاتورتكم رقم {invoice_no}\n\n{msg}"
                encoded = urllib.parse.quote(share_msg)
                try:
                    if phone:
                        clean_phone = ''.join(c for c in phone if c.isdigit())
                        webbrowser.open(f"whatsapp://send?phone={clean_phone}&text={encoded}")
                    else:
                        webbrowser.open(f"whatsapp://send?text={encoded}")
                except Exception:
                    pass

            btn_open.bind(on_press=open_pdf)
            btn_share.bind(on_press=share_wa)
            btn_close.bind(on_press=lambda x: popup2.dismiss())

            btn_row.add_widget(btn_close)
            btn_row.add_widget(btn_share)
            btn_row.add_widget(btn_open)
            box2.add_widget(btn_row)
            popup2.open()

        from kivy.clock import Clock
        Clock.schedule_once(do_generate, 0.3)

    def delete_shipment(self, shp_no):
        lbl_kwargs = {'font_name': ARABIC_FONT} if ARABIC_FONT else {}
        title_kwargs = {'font_name': ARABIC_FONT_BOLD} if ARABIC_FONT_BOLD else {}

        box = BoxLayout(orientation='vertical', padding=15, spacing=10)
        box.add_widget(Label(
            text=ar(f"حذف الشحنة رقم ({shp_no})؟"),
            font_size='15sp',
            halign='center',
            size_hint_y=None, height=30,
            **title_kwargs))

        box.add_widget(Label(
            text=ar("سيتم حذف القيود المحاسبية المرتبطة"),
            font_size='12sp',
            halign='center',
            color=(0.9, 0.6, 0.3, 1),
            size_hint_y=None, height=26,
            **lbl_kwargs))

        btn_row = BoxLayout(size_hint_y=None, height=45, spacing=8)
        btn_yes = Button(text=ar("حذف"), background_color=(0.7, 0.2, 0.2, 1),
                         background_normal='', **lbl_kwargs)
        btn_no = Button(text=ar("إلغاء"), background_color=(0.3, 0.35, 0.45, 1),
                        background_normal='', **lbl_kwargs)

        popup = Popup(
            title=ar("تأكيد الحذف"),
            content=box,
            size_hint=(0.85, 0.35),
            title_font=ARABIC_FONT_BOLD if ARABIC_FONT_BOLD else 'Roboto',
        )

        def confirm_del(inst):
            conn = get_conn()
            cursor = conn.cursor()
            cursor.execute("DELETE FROM shipments WHERE shipment_no=?", (shp_no,))
            cursor.execute("DELETE FROM shipment_custom_fees WHERE shipment_no=?", (shp_no,))
            cursor.execute("DELETE FROM journal_entries WHERE ref_no=?", (shp_no,))
            conn.commit()
            conn.close()
            popup.dismiss()
            self.load_data()

        btn_yes.bind(on_press=confirm_del)
        btn_no.bind(on_press=lambda x: popup.dismiss())

        btn_row.add_widget(btn_yes)
        btn_row.add_widget(btn_no)
        box.add_widget(btn_row)
        popup.open()


# ============================================================
# 14. شاشة إدارة الفواتير
# ============================================================
class InvoicesScreen(Screen):
    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.layout = BoxLayout(orientation='vertical', padding=(10, 12, 10, 8), spacing=6)
        self.add_widget(self.layout)

    def load_invoices(self):
        lbl_kwargs = {'font_name': ARABIC_FONT} if ARABIC_FONT else {}
        title_kwargs = {'font_name': ARABIC_FONT_BOLD} if ARABIC_FONT_BOLD else {}
        self.layout.clear_widgets()

        title_box = BoxLayout(size_hint_y=None, height=40)
        title_box.add_widget(Label(
            text=ar("سجل الفواتير"),
            font_size='17sp',
            color=(0.2, 0.8, 1, 1),
            **title_kwargs))
        self.layout.add_widget(title_box)

        scroll = ScrollView(size_hint=(1, 1))
        list_box = BoxLayout(
            orientation='vertical',
            size_hint_y=None,
            spacing=10,
            padding=(2, 4))
        list_box.bind(minimum_height=list_box.setter('height'))

        invoices = get_all_invoices()

        if not invoices:
            empty_box = BoxLayout(orientation='vertical', size_hint_y=None, height=100)
            empty_box.add_widget(Label(
                text=ar("لا توجد فواتير صادرة"),
                font_size='14sp',
                color=(0.6, 0.65, 0.75, 1),
                **lbl_kwargs))
            empty_box.add_widget(Label(
                text=ar("اذهب لسجل الشحنات واضغط على 'فاتورة'"),
                font_size='12sp',
                color=(0.5, 0.55, 0.65, 1),
                **lbl_kwargs))
            list_box.add_widget(empty_box)
        else:
            for inv in invoices:
                inv_id, inv_no, shp_no, client_name, amount, currency, pdf_path, notes, created = inv
                card = BoxLayout(orientation='vertical', size_hint_y=None, height=120,
                                 spacing=4, padding=(12, 8, 12, 8))

                with card.canvas.before:
                    Color(0.09, 0.12, 0.17, 1)
                    _cbg = RoundedRectangle(pos=card.pos, size=card.size, radius=[8])
                card.bind(
                    pos=lambda inst, val, bg=_cbg: setattr(bg, 'pos', val),
                    size=lambda inst, val, bg=_cbg: setattr(bg, 'size', val))

                row1 = BoxLayout(size_hint_y=None, height=28)
                lbl_no = Label(
                    text=ar(f"فاتورة رقم: {inv_no}"),
                    font_size='14sp',
                    color=(0.3, 0.85, 1, 1),
                    halign='left', valign='middle',
                    size_hint_x=0.5,
                    **title_kwargs)
                lbl_no.bind(size=lambda inst, val: setattr(inst, 'text_size', val))

                lbl_date = Label(
                    text=ar(f"{created or ''}"),
                    font_size='11sp',
                    color=(0.6, 0.7, 0.8, 1),
                    halign='right', valign='middle',
                    size_hint_x=0.5,
                    **lbl_kwargs)
                lbl_date.bind(size=lambda inst, val: setattr(inst, 'text_size', val))

                row1.add_widget(lbl_no)
                row1.add_widget(lbl_date)
                card.add_widget(row1)

                row2 = BoxLayout(size_hint_y=None, height=24)
                lbl_info = Label(
                    text=ar(f"العميل: {client_name}  |  الشحنة: {shp_no}  |  المبلغ: {fmt_num(amount)} {currency}"),
                    font_size='11sp',
                    color=(0.8, 0.8, 0.9, 1),
                    halign='right', valign='middle',
                    **lbl_kwargs)
                lbl_info.bind(size=lambda inst, val: setattr(inst, 'text_size', val))
                row2.add_widget(lbl_info)
                card.add_widget(row2)

                row3 = BoxLayout(size_hint_y=None, height=36, spacing=6)

                btn_open = Button(
                    text=ar("فتح PDF"),
                    font_size='11sp',
                    background_color=(0.15, 0.45, 0.8, 1),
                    background_normal='',
                    **lbl_kwargs)
                btn_open.bind(on_press=lambda x, p=pdf_path: self.open_pdf(p))

                btn_send = Button(
                    text=ar("إرسال واتساب"),
                    font_size='11sp',
                    background_color=(0, 0.65, 0.35, 1),
                    background_normal='',
                    **lbl_kwargs)
                btn_send.bind(on_press=lambda x, s=shp_no, i=inv_no: self.send_invoice_wa(s, i))

                btn_del = Button(
                    text=ar("حذف"),
                    font_size='11sp',
                    background_color=(0.75, 0.2, 0.2, 1),
                    background_normal='',
                    **lbl_kwargs)
                btn_del.bind(on_press=lambda x, i=inv_id: self.delete_inv(i))

                row3.add_widget(btn_open)
                row3.add_widget(btn_send)
                row3.add_widget(btn_del)
                card.add_widget(row3)

                list_box.add_widget(card)

        scroll.add_widget(list_box)
        self.layout.add_widget(scroll)

        btns = BoxLayout(size_hint_y=None, height=45, spacing=6)
        btn_back = Button(
            text=ar("الرئيسية"),
            background_color=(0.35, 0.4, 0.5, 1),
            background_normal='',
            font_size='13sp',
            **lbl_kwargs)
        btn_back.bind(on_press=lambda x: setattr(self.manager, 'current', 'home'))
        btns.add_widget(btn_back)
        self.layout.add_widget(btns)

    def open_pdf(self, pdf_path):
        try:
            if os.name == 'nt':
                os.startfile(pdf_path)
            else:
                webbrowser.open(f"file://{os.path.abspath(pdf_path)}")
        except Exception:
            pass

    def send_invoice_wa(self, shp_no, inv_no):
        msg, phone = build_whatsapp_msg(shp_no, include_phone_target=True)
        full_msg = f"*فاتورة رقم:* {inv_no}\n\n{msg}"
        encoded = urllib.parse.quote(full_msg)
        try:
            if phone:
                clean_phone = ''.join(c for c in phone if c.isdigit())
                webbrowser.open(f"whatsapp://send?phone={clean_phone}&text={encoded}")
            else:
                webbrowser.open(f"whatsapp://send?text={encoded}")
        except Exception:
            pass

    def delete_inv(self, inv_id):
        lbl_kwargs = {'font_name': ARABIC_FONT} if ARABIC_FONT else {}
        title_kwargs = {'font_name': ARABIC_FONT_BOLD} if ARABIC_FONT_BOLD else {}

        box = BoxLayout(orientation='vertical', padding=15, spacing=10)
        box.add_widget(Label(
            text=ar("حذف هذه الفاتورة؟"),
            font_size='15sp',
            halign='center',
            size_hint_y=None, height=30,
            **title_kwargs))

        btn_row = BoxLayout(size_hint_y=None, height=45, spacing=8)
        btn_yes = Button(text=ar("حذف"), background_color=(0.7, 0.2, 0.2, 1),
                         background_normal='', **lbl_kwargs)
        btn_no = Button(text=ar("إلغاء"), background_color=(0.3, 0.35, 0.45, 1),
                        background_normal='', **lbl_kwargs)

        popup = Popup(title=ar("تأكيد"), content=box, size_hint=(0.85, 0.3),
                      title_font=ARABIC_FONT_BOLD if ARABIC_FONT_BOLD else 'Roboto')

        def confirm(inst):
            delete_invoice_by_id(inv_id)
            popup.dismiss()
            self.load_invoices()

        btn_yes.bind(on_press=confirm)
        btn_no.bind(on_press=lambda x: popup.dismiss())

        btn_row.add_widget(btn_yes)
        btn_row.add_widget(btn_no)
        box.add_widget(btn_row)
        popup.open()


# ============================================================
# 15. شاشة كشف حساب عميل
# ============================================================
class CustomerStatementScreen(Screen):
    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.layout = BoxLayout(orientation='vertical', padding=(10, 12, 10, 8), spacing=6)
        self.add_widget(self.layout)
        self.current_client = None
        self.rows_cache = []
        self.date_from = None
        self.date_to = None

    def load_clients(self):
        lbl_kwargs = {'font_name': ARABIC_FONT} if ARABIC_FONT else {}
        title_kwargs = {'font_name': ARABIC_FONT_BOLD} if ARABIC_FONT_BOLD else {}
        self.layout.clear_widgets()
        self.current_client = None
        self.date_from = None
        self.date_to = None

        title_box = BoxLayout(size_hint_y=None, height=38)
        title_box.add_widget(Label(
            text=ar("كشف حساب عميل"),
            font_size='17sp',
            color=(0.2, 0.8, 1, 1),
            **title_kwargs))
        self.layout.add_widget(title_box)

        clients = get_all_clients()

        select_box = BoxLayout(size_hint_y=None, height=46, spacing=8)
        select_box.add_widget(Label(
            text=ar("اختر العميل:"),
            font_size='13sp',
            size_hint_x=0.3,
            halign='right',
            valign='middle',
            **lbl_kwargs))

        if clients:
            self.spn_client = Spinner(
                text=ar(clients[0]),
                values=[ar(c) for c in clients],
                option_cls=ArabicSpinnerOption,
                font_size='13sp',
                size_hint_x=0.7,
                **lbl_kwargs)
            self.spn_client.bind(text=lambda s, t: self.apply_filter())
        else:
            self.spn_client = Label(
                text=ar("لا يوجد عملاء مسجلون"),
                font_size='13sp',
                size_hint_x=0.7,
                color=(0.7, 0.7, 0.7, 1),
                **lbl_kwargs)

        select_box.add_widget(self.spn_client)
        self.layout.add_widget(select_box)

        filter_box = BoxLayout(size_hint_y=None, height=42, spacing=6, padding=(2, 2))

        btn_apply = Button(
            text=ar("تطبيق"),
            background_color=(0.15, 0.55, 0.75, 1),
            background_normal='',
            font_size='12sp',
            size_hint_x=0.16,
            **lbl_kwargs)
        btn_apply.bind(on_press=lambda x: self.apply_filter())
        filter_box.add_widget(btn_apply)

        self.inp_to = ArabicTextInput(
            hint_text="إلى YYYY-MM-DD",
            size_hint_x=0.32,
            size_hint_y=None,
            height=42)
        filter_box.add_widget(self.inp_to)

        filter_box.add_widget(Label(
            text=ar("إلى:"),
            font_size='12sp',
            size_hint_x=0.10,
            halign='right',
            valign='middle',
            **lbl_kwargs))

        self.inp_from = ArabicTextInput(
            hint_text="من YYYY-MM-DD",
            size_hint_x=0.32,
            size_hint_y=None,
            height=42)
        filter_box.add_widget(self.inp_from)

        filter_box.add_widget(Label(
            text=ar("من:"),
            font_size='12sp',
            size_hint_x=0.10,
            halign='right',
            valign='middle',
            **lbl_kwargs))

        self.layout.add_widget(filter_box)

        self.result_scroll = ScrollView(size_hint=(1, 1))
        self.result_box = BoxLayout(
            orientation='vertical',
            size_hint_y=None,
            spacing=8,
            padding=(2, 6))
        self.result_box.bind(minimum_height=self.result_box.setter('height'))
        self.result_scroll.add_widget(self.result_box)
        self.layout.add_widget(self.result_scroll)

        btns = BoxLayout(size_hint_y=None, height=44, spacing=6)

        btn_clear = Button(
            text=ar("مسح الفلتر"),
            background_color=(0.5, 0.4, 0.15, 1),
            background_normal='',
            font_size='12sp',
            **lbl_kwargs)
        btn_clear.bind(on_press=self.clear_filter)

        btn_export = Button(
            text=ar("تصدير Excel"),
            background_color=(0.1, 0.45, 0.3, 1),
            background_normal='',
            font_size='12sp',
            **lbl_kwargs)
        btn_export.bind(on_press=self.export_excel)

        btn_back = Button(
            text=ar("الرئيسية"),
            background_color=(0.35, 0.4, 0.5, 1),
            background_normal='',
            font_size='12sp',
            **lbl_kwargs)
        btn_back.bind(on_press=lambda x: setattr(self.manager, 'current', 'home'))

        btns.add_widget(btn_back)
        btns.add_widget(btn_clear)
        btns.add_widget(btn_export)
        self.layout.add_widget(btns)

        if clients:
            self.apply_filter()

    def _validate_dates(self, date_from, date_to):
        if date_from and date_to:
            if date_from > date_to:
                lbl_kwargs = {'font_name': ARABIC_FONT} if ARABIC_FONT else {}
                Popup(
                    title=ar("خطأ في التواريخ"),
                    content=Label(
                        text=ar("التاريخ (من) يجب أن يكون قبل التاريخ (إلى)"),
                        halign='center',
                        **lbl_kwargs),
                    size_hint=(0.85, 0.25)
                ).open()
                return False
        return True

    def clear_filter(self, instance):
        self.inp_from.text = ""
        self.inp_to.text = ""
        self.apply_filter()

    def apply_filter(self):
        if not hasattr(self, 'spn_client'):
            return
        current = self.spn_client.text
        clients = get_all_clients()
        client_name = None
        for c in clients:
            if ar(c) == current:
                client_name = c
                break
        if not client_name:
            return

        self.current_client = client_name
        date_from = self.inp_from.text.strip() or None
        date_to = self.inp_to.text.strip() or None

        if not self._validate_dates(date_from, date_to):
            return

        self.date_from = date_from
        self.date_to = date_to

        rows = get_client_statement(client_name, date_from, date_to)
        self.rows_cache = rows
        self._render_statement(rows)

    def _render_statement(self, rows):
        lbl_kwargs = {'font_name': ARABIC_FONT} if ARABIC_FONT else {}
        title_kwargs = {'font_name': ARABIC_FONT_BOLD} if ARABIC_FONT_BOLD else {}

        self.result_box.clear_widgets()

        if not rows:
            self.result_box.add_widget(Label(
                text=ar("لا توجد شحنات لهذا العميل في الفترة المحددة"),
                font_size='14sp',
                color=(0.7, 0.7, 0.7, 1),
                size_hint_y=None, height=60,
                **lbl_kwargs))
            return

        total_adv = 0.0
        total_exp = 0.0
        total_bal = 0.0
        for r in rows:
            total_adv += r[2] or 0
            total_exp += r[3] or 0
            total_bal += r[5] or 0

        summary = BoxLayout(
            orientation='vertical',
            size_hint_y=None,
            height=135,
            spacing=4,
            padding=(14, 10, 14, 10))

        with summary.canvas.before:
            Color(0.12, 0.16, 0.22, 1)
            _sbg = RoundedRectangle(pos=summary.pos, size=summary.size, radius=[10])
        summary.bind(
            pos=lambda inst, val, bg=_sbg: setattr(bg, 'pos', val),
            size=lambda inst, val, bg=_sbg: setattr(bg, 'size', val))

        period_str = ""
        if self.date_from or self.date_to:
            period_str = f"  (الفترة: {self.date_from or '...'} → {self.date_to or '...'})"

        summary.add_widget(Label(
            text=ar(f"عدد الشحنات: {len(rows)}{period_str}"),
            font_size='13sp',
            color=(0.3, 0.85, 1, 1),
            halign='right', valign='middle',
            size_hint_y=None, height=24,
            **title_kwargs))

        summary.add_widget(Label(
            text=ar(f"إجمالي العهد المقبوضة:  {fmt_num(total_adv)}"),
            font_size='13sp',
            color=(0.4, 0.7, 1, 1),
            halign='right', valign='middle',
            size_hint_y=None, height=22,
            **lbl_kwargs))

        summary.add_widget(Label(
            text=ar(f"إجمالي المصاريف والأتعاب:  {fmt_num(total_exp)}"),
            font_size='13sp',
            color=(1, 0.7, 0.3, 1),
            halign='right', valign='middle',
            size_hint_y=None, height=22,
            **lbl_kwargs))

        if total_bal >= 0:
            bal_color = (0.3, 0.9, 0.5, 1)
            bal_label = f"الرصيد النهائي (دائن للعميل):  {fmt_num(abs(total_bal))}"
        else:
            bal_color = (1, 0.4, 0.4, 1)
            bal_label = f"الرصيد النهائي (مدين على العميل):  {fmt_num(abs(total_bal))}"

        summary.add_widget(Label(
            text=ar(bal_label),
            font_size='13sp',
            color=bal_color,
            halign='right', valign='middle',
            size_hint_y=None, height=28,
            **title_kwargs))

        self.result_box.add_widget(summary)

        for r in rows:
            shp_no, curr, adv, exp, fee, bal, status, date = r

            card = BoxLayout(
                orientation='vertical',
                size_hint_y=None,
                height=165,
                spacing=5,
                padding=(14, 10, 14, 10))

            with card.canvas.before:
                Color(0.09, 0.12, 0.17, 1)
                _cbg = RoundedRectangle(pos=card.pos, size=card.size, radius=[8])
            card.bind(
                pos=lambda inst, val, bg=_cbg: setattr(bg, 'pos', val),
                size=lambda inst, val, bg=_cbg: setattr(bg, 'size', val))

            row1 = BoxLayout(size_hint_y=None, height=28, spacing=8)

            lbl_no = Label(
                text=ar(f"شحنة رقم: {shp_no}"),
                font_size='14sp',
                color=(0.3, 0.85, 1, 1),
                halign='left', valign='middle',
                size_hint_x=0.5,
                **title_kwargs)
            lbl_no.bind(size=lambda inst, val: setattr(inst, 'text_size', val))

            status_clean = "قيد الإجراء" if "تحت" in status else "مكتملة ومسلمة"
            status_color = (0.3, 0.9, 0.5, 1) if "مكتملة" in status else (1, 0.75, 0.3, 1)

            lbl_st = Label(
                text=ar(status_clean),
                font_size='12sp',
                color=status_color,
                halign='right', valign='middle',
                size_hint_x=0.5,
                **lbl_kwargs)
            lbl_st.bind(size=lambda inst, val: setattr(inst, 'text_size', val))

            row1.add_widget(lbl_no)
            row1.add_widget(lbl_st)
            card.add_widget(row1)

            sep = BoxLayout(size_hint_y=None, height=1)
            with sep.canvas.before:
                Color(0.2, 0.28, 0.38, 1)
                _sep2 = Rectangle(pos=sep.pos, size=sep.size)
            sep.bind(
                pos=lambda inst, val, s=_sep2: setattr(s, 'pos', val),
                size=lambda inst, val, s=_sep2: setattr(s, 'size', val))
            card.add_widget(sep)

            lbl_adv = Label(
                text=ar(f"العهد المقبوضة:  {fmt_num(adv)} {curr}"),
                font_size='12sp',
                color=(0.4, 0.7, 1, 1),
                halign='right', valign='middle',
                size_hint_y=None, height=24,
                **lbl_kwargs)
            lbl_adv.bind(size=lambda inst, val: setattr(inst, 'text_size', (val[0] - 4, val[1])))
            card.add_widget(lbl_adv)

            lbl_exp = Label(
                text=ar(f"المصاريف والأتعاب:  {fmt_num(exp)} {curr}"),
                font_size='12sp',
                color=(1, 0.7, 0.3, 1),
                halign='right', valign='middle',
                size_hint_y=None, height=24,
                **lbl_kwargs)
            lbl_exp.bind(size=lambda inst, val: setattr(inst, 'text_size', (val[0] - 4, val[1])))
            card.add_widget(lbl_exp)

            if bal >= 0:
                bc = (0.3, 0.9, 0.5, 1)
                bt = f"الرصيد (دائن):  {fmt_num(abs(bal))} {curr}"
            else:
                bc = (1, 0.4, 0.4, 1)
                bt = f"الرصيد (مدين):  {fmt_num(abs(bal))} {curr}"

            lbl_bal = Label(
                text=ar(bt),
                font_size='13sp',
                color=bc,
                halign='right', valign='middle',
                size_hint_y=None, height=28,
                **title_kwargs)
            lbl_bal.bind(size=lambda inst, val: setattr(inst, 'text_size', (val[0] - 4, val[1])))
            card.add_widget(lbl_bal)

            self.result_box.add_widget(card)

    def export_excel(self, instance):
        if not self.current_client:
            return
        try:
            filename = export_client_statement_excel(
                self.current_client, self.date_from, self.date_to)
            if filename:
                lbl_kwargs = {'font_name': ARABIC_FONT} if ARABIC_FONT else {}
                Popup(
                    title=ar("تم التصدير"),
                    content=Label(text=ar(f"تم الحفظ في:\n{filename}"), **lbl_kwargs),
                    size_hint=(0.85, 0.3)
                ).open()
        except Exception:
            pass


# ============================================================
# 16. شاشة المصروفات
# ============================================================
class ExpensesScreen(Screen):
    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.layout = BoxLayout(orientation='vertical', padding=10, spacing=8)
        self.add_widget(self.layout)

    def load_categories(self):
        lbl_kwargs = {'font_name': ARABIC_FONT} if ARABIC_FONT else {}
        title_kwargs = {'font_name': ARABIC_FONT_BOLD} if ARABIC_FONT_BOLD else {}
        self.layout.clear_widgets()

        self.layout.add_widget(Label(text=ar("إدارة المصروفات القابلة للتصنيف"),
                                     font_size='17sp', size_hint_y=0.08, **title_kwargs))

        form = GridLayout(cols=2, spacing=10, size_hint_y=0.6)

        form.add_widget(Label(text=ar("التاريخ:"), size_hint_x=0.38, **lbl_kwargs))
        self.inp_date = ArabicTextInput(text=datetime.now().strftime("%Y-%m-%d"),
                                        size_hint_x=0.62, size_hint_y=None, height=60)
        form.add_widget(self.inp_date)

        form.add_widget(Label(text=ar("بيان المصروف:"), size_hint_x=0.38, **lbl_kwargs))
        self.inp_title = ArabicTextInput(hint_text="مثال: فاتورة كهرباء / صيانة",
                                         size_hint_x=0.62, size_hint_y=None, height=60)
        form.add_widget(self.inp_title)

        form.add_widget(Label(text=ar("التصنيف:"), size_hint_x=0.38, **lbl_kwargs))
        conn = get_conn()
        cursor = conn.cursor()
        cursor.execute("SELECT cat_name FROM exp_categories")
        exp_cats = [ar(r[0]) for r in cursor.fetchall()]
        conn.close()
        exp_cats.append(ar("+ إضافة تصنيف جديد"))

        self.spn_cat = Spinner(
            text=exp_cats[0] if exp_cats else ar('تشغيلي'),
            values=exp_cats,
            option_cls=ArabicSpinnerOption,
            size_hint_x=0.62,
            size_hint_y=None,
            height=60,
            **lbl_kwargs)
        self.spn_cat.bind(text=self.on_category_select)
        form.add_widget(self.spn_cat)

        form.add_widget(Label(text=ar("المبلغ:"), size_hint_x=0.38, **lbl_kwargs))
        self.inp_amt = ArabicTextInput(text="0", input_filter='float',
                                       size_hint_x=0.62, size_hint_y=None, height=60)
        form.add_widget(self.inp_amt)

        form.add_widget(Label(text=ar("العملة:"), size_hint_x=0.38, **lbl_kwargs))
        self.spn_exp_curr = Spinner(text='YER', values=('YER', 'SAR', 'USD'),
                                    option_cls=ArabicSpinnerOption,
                                    size_hint_x=0.62, size_hint_y=None, height=60,
                                    **lbl_kwargs)
        form.add_widget(self.spn_exp_curr)

        self.layout.add_widget(form)
        self.lbl_exp_status = Label(text="", size_hint_y=0.08, **lbl_kwargs)
        self.layout.add_widget(self.lbl_exp_status)

        btns = BoxLayout(size_hint_y=0.14, spacing=10)
        btn_save = Button(text=ar("حفظ القيد المحاسبي"),
                          background_color=(0.8, 0.3, 0.1, 1),
                          background_normal='', **lbl_kwargs)
        btn_save.bind(on_press=self.save_expense)
        btn_back = Button(text=ar("رجوع"),
                          background_color=(0.35, 0.4, 0.5, 1),
                          background_normal='', **lbl_kwargs)
        btn_back.bind(on_press=lambda x: setattr(self.manager, 'current', 'home'))

        btns.add_widget(btn_save)
        btns.add_widget(btn_back)
        self.layout.add_widget(btns)

    def on_category_select(self, spinner, text):
        if text == ar("+ إضافة تصنيف جديد"):
            self.popup_add_category()

    def popup_add_category(self):
        lbl_kwargs = {'font_name': ARABIC_FONT} if ARABIC_FONT else {}
        box = BoxLayout(orientation='vertical', padding=12, spacing=10)
        box.add_widget(Label(text=ar("أدخل اسم التصنيف الجديد:"), **lbl_kwargs))

        inp_new_cat = ArabicTextInput(hint_text="اسم التصنيف", size_hint_y=None, height=60)
        box.add_widget(inp_new_cat)

        btn_box = BoxLayout(size_hint_y=None, height=50, spacing=5)
        btn_add = Button(text=ar("إضافة"),
                         background_color=(0, 0.6, 0.3, 1),
                         background_normal='', **lbl_kwargs)
        popup = Popup(title=ar("إضافة تصنيف مصروفات"), content=box, size_hint=(0.85, 0.4))

        def do_add(instance):
            cat_val = inp_new_cat.text.strip()
            if cat_val:
                conn = get_conn()
                cursor = conn.cursor()
                cursor.execute("INSERT OR IGNORE INTO exp_categories (cat_name) VALUES (?)", (cat_val,))
                conn.commit()
                conn.close()
                self.load_categories()
                if ar(cat_val) in self.spn_cat.values:
                    self.spn_cat.text = ar(cat_val)
            popup.dismiss()

        btn_add.bind(on_press=do_add)
        btn_box.add_widget(btn_add)
        box.add_widget(btn_box)
        popup.open()

    def save_expense(self, instance):
        try:
            d = self.inp_date.text.strip()
            t = self.inp_title.text.strip()
            c = self.spn_cat.text
            try:
                a = float(self.inp_amt.text or 0)
            except ValueError:
                a = 0.0
            curr = self.spn_exp_curr.text

            if not t:
                self.lbl_exp_status.text = ar("خطأ: بيان المصروف مطلوب")
                self.lbl_exp_status.color = (1, 0, 0, 1)
                return

            conn = get_conn()
            cursor = conn.cursor()
            cursor.execute(
                "INSERT INTO expenses (exp_date, exp_title, category, amount, currency) VALUES (?,?,?,?,?)",
                (d, t, c, a, curr))

            cursor.execute(
                "INSERT INTO journal_entries (entry_date, ref_no, account_code, debit, credit, currency, description) VALUES (?,?,?,?,?,?,?)",
                (d, 'EXP', '5101', a, 0, curr, f"مصروف: {t} - {c}"))
            cursor.execute(
                "INSERT INTO journal_entries (entry_date, ref_no, account_code, debit, credit, currency, description) VALUES (?,?,?,?,?,?,?)",
                (d, 'EXP', '1101', 0, a, curr, f"مصروف: {t} - {c}"))

            conn.commit()
            conn.close()
            self.lbl_exp_status.text = ar("تم تسجيل القيد المحاسبي للمصروف!")
            self.lbl_exp_status.color = (0, 1, 0, 1)
        except Exception as e:
            self.lbl_exp_status.text = ar(f"خطأ: {str(e)}")
            self.lbl_exp_status.color = (1, 0, 0, 1)


# ============================================================
# 17. شاشة التقارير المالية
# ============================================================
class DashboardScreen(Screen):
    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.layout = BoxLayout(orientation='vertical', padding=12, spacing=8)
        self.add_widget(self.layout)
        self.sub_container = None

    def load_main_menu(self):
        lbl_kwargs = {'font_name': ARABIC_FONT} if ARABIC_FONT else {}
        title_kwargs = {'font_name': ARABIC_FONT_BOLD} if ARABIC_FONT_BOLD else {}
        self.layout.clear_widgets()

        self.layout.add_widget(Label(
            text=ar("التقرير المالي والقوائم المالية"),
            font_size='16sp', size_hint_y=None, height=40,
            color=(0.3, 0.85, 1, 1), **title_kwargs))

        menu_btns = [
            ("قائمة الدخل (Income Statement)", self.show_income_statement, (0.1, 0.45, 0.75, 1)),
            ("الميزانية العمومية (Balance Sheet)", self.show_balance_sheet, (0.15, 0.55, 0.35, 1)),
            ("ميزان المراجعة (Trial Balance)", self.show_trial_balance, (0.65, 0.35, 0.15, 1)),
        ]
        for title, func, col in menu_btns:
            btn = Button(text=ar(title), font_size='13sp',
                         size_hint_y=None, height=48,
                         background_color=col,
                         background_normal='', **lbl_kwargs)
            btn.bind(on_press=func)
            self.layout.add_widget(btn)

        result_scroll = ScrollView(size_hint=(1, 1))
        self.sub_container = BoxLayout(
            orientation='vertical',
            size_hint_y=None, spacing=6, padding=6)
        self.sub_container.bind(minimum_height=self.sub_container.setter('height'))
        result_scroll.add_widget(self.sub_container)
        self.layout.add_widget(result_scroll)

        btn_back = Button(text=ar("الرئيسية"), size_hint_y=None, height=45,
                          background_color=(0.35, 0.4, 0.5, 1),
                          background_normal='', **lbl_kwargs)
        btn_back.bind(on_press=lambda x: setattr(self.manager, 'current', 'home'))
        self.layout.add_widget(btn_back)

    def _add_result_row(self, label_text, value, color=(1, 1, 1, 1), is_title=False):
        lbl_kwargs = {'font_name': ARABIC_FONT} if ARABIC_FONT else {}
        title_kwargs = {'font_name': ARABIC_FONT_BOLD} if ARABIC_FONT_BOLD else {}
        row = BoxLayout(size_hint_y=None, height=45, spacing=6)
        lbl = Label(
            text=ar(label_text),
            font_size='14sp' if is_title else '13sp',
            halign='right', valign='middle',
            size_hint_x=0.62,
            color=(0.3, 0.85, 1, 1) if is_title else (0.9, 0.9, 0.9, 1),
            **(title_kwargs if is_title else lbl_kwargs))
        lbl.bind(size=lambda inst, val: setattr(inst, 'text_size', val))
        val_lbl = Label(
            text=str(value),
            font_size='14sp' if is_title else '13sp',
            halign='left', valign='middle',
            size_hint_x=0.38, color=color,
            **lbl_kwargs)
        val_lbl.bind(size=lambda inst, val: setattr(inst, 'text_size', val))
        row.add_widget(lbl)
        row.add_widget(val_lbl)
        self.sub_container.add_widget(row)

    def _get_balances(self):
        conn = get_conn()
        cursor = conn.cursor()
        cursor.execute("""SELECT account_code, SUM(debit), SUM(credit)
                          FROM journal_entries GROUP BY account_code""")
        rows = cursor.fetchall()
        conn.close()
        return {code: {'debit': d or 0, 'credit': c or 0} for code, d, c in rows}

    def show_income_statement(self, instance):
        self.sub_container.clear_widgets()
        balances = self._get_balances()
        revenue = balances.get('4101', {}).get('credit', 0) - balances.get('4101', {}).get('debit', 0)
        expenses = balances.get('5101', {}).get('debit', 0) - balances.get('5101', {}).get('credit', 0)
        net = revenue - expenses

        self._add_result_row("— قائمة الدخل —", "", (1, 0.8, 0.2, 1), is_title=True)
        self._add_result_row("إجمالي إيراد الأتعاب:", fmt_num(revenue), (0, 1, 0, 1))
        self._add_result_row("المصروفات العمومية والتشغيلية:", fmt_num(expenses), (1, 0.4, 0.4, 1))
        net_color = (0, 1, 0, 1) if net >= 0 else (1, 0.3, 0.3, 1)
        self._add_result_row("صافي الربح / الخسارة:", fmt_num(net), net_color)

    def show_balance_sheet(self, instance):
        self.sub_container.clear_widgets()
        balances = self._get_balances()
        assets_1101 = balances.get('1101', {}).get('debit', 0) - balances.get('1101', {}).get('credit', 0)
        liab_2101 = balances.get('2101', {}).get('credit', 0) - balances.get('2101', {}).get('debit', 0)
        revenue = balances.get('4101', {}).get('credit', 0) - balances.get('4101', {}).get('debit', 0)
        expenses = balances.get('5101', {}).get('debit', 0) - balances.get('5101', {}).get('credit', 0)
        equity = revenue - expenses
        check = assets_1101 - (liab_2101 + equity)

        self._add_result_row("— الميزانية العمومية —", "", (1, 0.8, 0.2, 1), is_title=True)
        self._add_result_row("النقدية والبنوك (أصول):", fmt_num(assets_1101), (0.2, 0.8, 1, 1))
        self._add_result_row("أمانات عهد العملاء (التزامات):", fmt_num(liab_2101), (1, 0.6, 0.2, 1))
        self._add_result_row("حقوق الملكية (أرباح مبقاة):", fmt_num(equity), (0, 1, 0, 1))
        check_color = (0, 1, 0, 1) if abs(check) < 0.01 else (1, 0.3, 0.3, 1)
        self._add_result_row("فرق الميزانية:", fmt_num(check), check_color)

    def show_trial_balance(self, instance):
        self.sub_container.clear_widgets()
        conn = get_conn()
        cursor = conn.cursor()
        cursor.execute("SELECT SUM(debit), SUM(credit) FROM journal_entries")
        row = cursor.fetchone()
        deb = row[0] or 0.0
        crd = row[1] or 0.0
        conn.close()
        balanced = abs(deb - crd) < 0.01

        self._add_result_row("— ميزان المراجعة —", "", (1, 0.8, 0.2, 1), is_title=True)
        self._add_result_row("إجمالي الجانب المدين (Debit):", fmt_num(deb), (0.2, 0.8, 1, 1))
        self._add_result_row("إجمالي الجانب الدائن (Credit):", fmt_num(crd), (0.2, 0.8, 1, 1))
        self._add_result_row("حالة التوازن المحاسبي:",
                             ar("متوازن 100%") if balanced else ar("غير متوازن"),
                             (0, 1, 0, 1) if balanced else (1, 0.3, 0.3, 1))


# ============================================================
# 18. شاشة إضافة/تعديل شحنة
# ============================================================
class AddShipmentScreen(Screen):
    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        layout = BoxLayout(orientation='vertical', padding=10, spacing=6)
        lbl_kwargs = {'font_name': ARABIC_FONT} if ARABIC_FONT else {}
        title_kwargs = {'font_name': ARABIC_FONT_BOLD} if ARABIC_FONT_BOLD else {}

        top_bar = BoxLayout(size_hint_y=0.07, spacing=10)
        self.title_lbl = Label(text=ar("إدخال البيانات المالية للشحنة الجمركية"),
                               font_size='15sp', color=(0.2, 0.8, 1, 1), **title_kwargs)
        top_bar.add_widget(self.title_lbl)

        btn_add_fee_type = Button(text=ar("+ إضافة بند رسوم"), size_hint_x=0.4,
                                  background_color=(0.2, 0.5, 0.8, 1),
                                  background_normal='',
                                  font_size='12sp', **lbl_kwargs)
        btn_add_fee_type.bind(on_press=self.popup_add_fee_type)
        top_bar.add_widget(btn_add_fee_type)
        layout.add_widget(top_bar)

        scroll = ScrollView(size_hint_y=0.72)
        self.form = GridLayout(cols=2, spacing=12, size_hint_y=None)
        self.form.bind(minimum_height=self.form.setter('height'))

        self.inputs = {}
        self.custom_fee_inputs = {}
        self.spn_curr = None
        self.spn_status = None
        self.spn_client = None

        self.build_form_fields()

        scroll.add_widget(self.form)
        layout.add_widget(scroll)

        self.lbl_status = Label(text="", size_hint_y=0.04, **lbl_kwargs)
        layout.add_widget(self.lbl_status)

        btns = BoxLayout(size_hint_y=0.08, spacing=10)
        self.btn_save_action = Button(text=ar("حفظ وترحيل قيود اليومية"), font_size='14sp',
                                      background_color=(0, 0.6, 0.3, 1),
                                      background_normal='', **lbl_kwargs)
        self.btn_save_action.bind(on_press=self.save_shipment)

        btn_back = Button(text=ar("رجوع"), font_size='14sp',
                          background_color=(0.35, 0.4, 0.5, 1),
                          background_normal='', **lbl_kwargs)
        btn_back.bind(on_press=self.go_home)

        btns.add_widget(self.btn_save_action)
        btns.add_widget(btn_back)
        layout.add_widget(btns)

        self.add_widget(layout)

    def go_home(self, instance):
        self.reset_form_mode()
        self.manager.current = 'home'

    def load_shipment_data_for_edit(self, shp_no):
        self.title_lbl.text = ar(f"تعديل الشحنة رقم: {shp_no}")
        self.btn_save_action.text = ar("حفظ التعديلات وترحيل القيود")

        conn = get_conn()
        cursor = conn.cursor()
        cursor.execute("""SELECT shipment_no, client_name, currency, advance_amount, customs_duty,
                          storage_fee, transport_fee, quality_fee, agriculture_fee, insurance_fee,
                          other_fee_amount, clearance_fee, doc_path, status
                          FROM shipments WHERE shipment_no=?""", (shp_no,))
        row = cursor.fetchone()

        cursor.execute("SELECT fee_name, amount FROM shipment_custom_fees WHERE shipment_no=?", (shp_no,))
        custom_fees_db = dict(cursor.fetchall())
        conn.close()

        if row:
            self.inputs['inp_no'].text = str(row[0])
            self.inputs['inp_no'].real_input.disabled = True
            try:
                self.spn_client.text = ar(str(row[1]))
            except Exception:
                pass
            self.spn_curr.text = str(row[2])
            self.inputs['inp_adv'].text = str(row[3])
            self.inputs['inp_duty'].text = str(row[4])
            self.inputs['inp_storage'].text = str(row[5])
            self.inputs['inp_transport'].text = str(row[6])
            self.inputs['inp_quality'].text = str(row[7])
            self.inputs['inp_agri'].text = str(row[8])
            self.inputs['inp_ins'].text = str(row[9])
            self.inputs['inp_other_amt'].text = str(row[10])
            self.inputs['inp_fee'].text = str(row[11])
            self.inputs['inp_doc'].text = str(row[12] or '')
            self.spn_status.text = ar(str(row[13]))

            for f_name, inp_obj in self.custom_fee_inputs.items():
                inp_obj.text = str(custom_fees_db.get(f_name, 0))

    def reset_form_mode(self):
        self.title_lbl.text = ar("إدخال البيانات المالية للشحنة الجمركية")
        self.btn_save_action.text = ar("حفظ وترحيل قيود اليومية")
        self.lbl_status.text = ""
        self.build_form_fields()

    def _clear_form(self):
        self.form.clear_widgets()
        self.inputs = {}
        self.custom_fee_inputs = {}

    def build_form_fields(self):
        self._clear_form()
        lbl_kwargs = {'font_name': ARABIC_FONT} if ARABIC_FONT else {}

        lbl = Label(text=ar("رقم الشحنة:"), font_size='13sp', halign='right',
                    valign='middle', size_hint_x=0.38, size_hint_y=None, height=66, **lbl_kwargs)
        lbl.bind(size=lambda inst, val: setattr(inst, 'text_size', val))
        self.form.add_widget(lbl)

        inp_no = ArabicTextInput(text="", input_filter=None,
                                  size_hint_x=0.62, size_hint_y=None, height=66)
        self.inputs['inp_no'] = inp_no
        self.form.add_widget(inp_no)

        lbl = Label(text=ar("اسم العميل:"), font_size='13sp', halign='right',
                    valign='middle', size_hint_x=0.38, size_hint_y=None, height=66, **lbl_kwargs)
        lbl.bind(size=lambda inst, val: setattr(inst, 'text_size', val))
        self.form.add_widget(lbl)

        clients_db = get_all_clients_objects()
        client_names_ar = [ar(c[1]) for c in clients_db]
        client_names_ar.append(ar("+ عميل جديد"))

        self.spn_client = Spinner(
            text=client_names_ar[0] if client_names_ar else ar("اختر عميلاً"),
            values=client_names_ar,
            option_cls=ArabicSpinnerOption,
            font_size='13sp',
            size_hint_x=0.62,
            size_hint_y=None,
            height=66,
            **lbl_kwargs)
        self.spn_client.bind(text=self.on_client_selected)
        self.form.add_widget(self.spn_client)

        fields = [
            ("العهدة المقبوضة (أمانة):", "inp_adv", True),
            ("الرسوم الجمركية الأساسية:", "inp_duty", True),
            ("رسوم الأرضيات والغرامات:", "inp_storage", True),
            ("مصاريف النقل والتحسين:", "inp_transport", True),
            ("رسوم الجودة:", "inp_quality", True),
            ("رسوم الزراعة:", "inp_agri", True),
            ("رسوم التأمين:", "inp_ins", True),
            ("مصاريف ورسوم أخرى:", "inp_other_amt", True),
            ("أتعاب التخليص الجمركي:", "inp_fee", True),
            ("رقم البيان/المستند:", "inp_doc", False),
        ]

        for label_text, name, is_numeric in fields:
            lbl = Label(text=ar(label_text), font_size='13sp', halign='right',
                        valign='middle', size_hint_x=0.38, size_hint_y=None, height=66, **lbl_kwargs)
            lbl.bind(size=lambda inst, val: setattr(inst, 'text_size', val))
            self.form.add_widget(lbl)

            default_txt = "0" if is_numeric else ""
            filt = 'float' if is_numeric else None

            inp = ArabicTextInput(text=default_txt, input_filter=filt,
                                  size_hint_x=0.62, size_hint_y=None, height=66)
            self.inputs[name] = inp
            self.form.add_widget(inp)

        conn = get_conn()
        cursor = conn.cursor()
        cursor.execute("SELECT fee_name FROM fee_types")
        fee_rows = [r[0] for r in cursor.fetchall()]
        conn.close()

        built_in_fees = {"الرسوم الجمركية الأساسية", "رسوم الأرضيات والغرامات",
                         "مصاريف النقل والتحسين", "رسوم الجودة", "رسوم الزراعة",
                         "رسوم التأمين", "أتعاب التخليص الجمركي"}

        for f_name in fee_rows:
            if f_name not in built_in_fees:
                lbl = Label(text=ar(f"{f_name}:"), font_size='13sp', halign='right',
                            valign='middle', size_hint_x=0.38, size_hint_y=None,
                            height=66, **lbl_kwargs)
                lbl.bind(size=lambda inst, val: setattr(inst, 'text_size', val))
                self.form.add_widget(lbl)

                inp = ArabicTextInput(text="0", input_filter='float',
                                      size_hint_x=0.62, size_hint_y=None, height=66)
                self.custom_fee_inputs[f_name] = inp
                self.form.add_widget(inp)

        self.form.add_widget(Label(text=ar("العملة:"), font_size='13sp',
                                   size_hint_x=0.38, size_hint_y=None, height=58, **lbl_kwargs))
        self.spn_curr = Spinner(text='YER', values=('YER', 'SAR', 'USD'),
                                option_cls=ArabicSpinnerOption,
                                size_hint_x=0.62, size_hint_y=None, height=58,
                                **lbl_kwargs)
        self.form.add_widget(self.spn_curr)

        self.form.add_widget(Label(text=ar("حالة الشحنة:"), font_size='13sp',
                                   size_hint_x=0.38, size_hint_y=None, height=58, **lbl_kwargs))
        self.spn_status = Spinner(text=ar('تحت الإجراء'),
                                  values=(ar('تحت الإجراء'), ar('مكتملة ومسلمة')),
                                  option_cls=ArabicSpinnerOption,
                                  size_hint_x=0.62, size_hint_y=None, height=58,
                                  **lbl_kwargs)
        self.form.add_widget(self.spn_status)

    def on_client_selected(self, spinner, text):
        if text == ar("+ عميل جديد"):
            self.open_new_client_popup()

    def open_new_client_popup(self):
        lbl_kwargs = {'font_name': ARABIC_FONT} if ARABIC_FONT else {}
        title_kwargs = {'font_name': ARABIC_FONT_BOLD} if ARABIC_FONT_BOLD else {}

        box = BoxLayout(orientation='vertical', padding=12, spacing=8)
        box.add_widget(Label(
            text=ar("إضافة عميل جديد"),
            font_size='14sp', size_hint_y=None, height=28,
            color=(0.3, 0.85, 1, 1),
            **title_kwargs))

        inp_name = ArabicTextInput(hint_text="اسم العميل *", size_hint_y=None, height=48)
        box.add_widget(inp_name)

        inp_phone = ArabicTextInput(hint_text="رقم الهاتف (اختياري)", size_hint_y=None, height=48)
        box.add_widget(inp_phone)

        btn_row = BoxLayout(size_hint_y=None, height=45, spacing=6)
        btn_save = Button(text=ar("حفظ"), background_color=(0, 0.6, 0.3, 1),
                          background_normal='', **lbl_kwargs)
        btn_cancel = Button(text=ar("إلغاء"), background_color=(0.4, 0.4, 0.5, 1),
                            background_normal='', **lbl_kwargs)

        popup = Popup(title=ar("عميل جديد"), content=box, size_hint=(0.9, 0.4),
                      title_font=ARABIC_FONT_BOLD if ARABIC_FONT_BOLD else 'Roboto')

        def do_save(inst):
            name = inp_name.text.strip()
            if not name:
                return
            phone = inp_phone.text.strip()
            ok, msg = save_client(None, name, phone, '', '', '', 0.0, '')
            if ok:
                popup.dismiss()
                self.build_form_fields()
                try:
                    self.spn_client.text = ar(name)
                except Exception:
                    pass

        btn_save.bind(on_press=do_save)
        btn_cancel.bind(on_press=lambda x: popup.dismiss())

        btn_row.add_widget(btn_cancel)
        btn_row.add_widget(btn_save)
        box.add_widget(btn_row)
        popup.open()

    def popup_add_fee_type(self, instance):
        lbl_kwargs = {'font_name': ARABIC_FONT} if ARABIC_FONT else {}
        box = BoxLayout(orientation='vertical', padding=12, spacing=10)
        box.add_widget(Label(text=ar("أدخل اسم بند الرسوم الجديد:"), **lbl_kwargs))

        inp_new = ArabicTextInput(hint_text="اسم البند", size_hint_y=None, height=60)
        box.add_widget(inp_new)

        btn_box = BoxLayout(size_hint_y=None, height=50, spacing=10)
        btn_add = Button(text=ar("إضافة"),
                         background_color=(0, 0.6, 0.3, 1),
                         background_normal='', **lbl_kwargs)
        popup = Popup(title=ar("إضافة رسوم جديدة"), content=box, size_hint=(0.85, 0.4))

        current_values = {k: v.text for k, v in self.inputs.items()}
        current_custom = {k: v.text for k, v in self.custom_fee_inputs.items()}
        current_curr = self.spn_curr.text if self.spn_curr else 'YER'
        current_status = self.spn_status.text if self.spn_status else ar('تحت الإجراء')
        current_client = self.spn_client.text if self.spn_client else ''

        def do_add(inst):
            val = inp_new.text.strip()
            if val:
                conn = get_conn()
                cursor = conn.cursor()
                cursor.execute("INSERT OR IGNORE INTO fee_types (fee_name) VALUES (?)", (val,))
                conn.commit()
                conn.close()
                self.build_form_fields()
                for k, v in current_values.items():
                    if k in self.inputs:
                        self.inputs[k].text = v
                for k, v in current_custom.items():
                    if k in self.custom_fee_inputs:
                        self.custom_fee_inputs[k].text = v
                if self.spn_curr:
                    self.spn_curr.text = current_curr
                if self.spn_status:
                    self.spn_status.text = current_status
                if self.spn_client:
                    self.spn_client.text = current_client
            popup.dismiss()

        btn_add.bind(on_press=do_add)
        btn_box.add_widget(btn_add)
        box.add_widget(btn_box)
        popup.open()

    def _safe_float(self, text):
        try:
            return float(text or 0)
        except (ValueError, TypeError):
            return 0.0

    def save_shipment(self, instance):
        try:
            shp_no = self.inputs['inp_no'].text.strip()
            client_ar = self.spn_client.text

            if not shp_no:
                self.lbl_status.text = ar("خطأ: رقم الشحنة مطلوب")
                self.lbl_status.color = (1, 0, 0, 1)
                return
            if not client_ar or client_ar == ar("+ عميل جديد"):
                self.lbl_status.text = ar("خطأ: الرجاء اختيار عميل")
                self.lbl_status.color = (1, 0, 0, 1)
                return

            client = None
            client_id = None
            for c in get_all_clients_objects():
                if ar(c[1]) == client_ar:
                    client = c[1]
                    client_id = c[0]
                    break

            if not client:
                self.lbl_status.text = ar("خطأ: العميل غير موجود")
                self.lbl_status.color = (1, 0, 0, 1)
                return

            adv = self._safe_float(self.inputs['inp_adv'].text)
            duty = self._safe_float(self.inputs['inp_duty'].text)
            storage = self._safe_float(self.inputs['inp_storage'].text)
            transport = self._safe_float(self.inputs['inp_transport'].text)
            quality = self._safe_float(self.inputs['inp_quality'].text)
            agri = self._safe_float(self.inputs['inp_agri'].text)
            ins = self._safe_float(self.inputs['inp_ins'].text)
            other_amt = self._safe_float(self.inputs['inp_other_amt'].text)
            fee = self._safe_float(self.inputs['inp_fee'].text)
            doc = self.inputs['inp_doc'].text.strip()
            curr = self.spn_curr.text
            status = self.spn_status.text
            created_at = datetime.now().strftime("%Y-%m-%d")

            custom_total = 0.0
            conn = get_conn()
            cursor = conn.cursor()
            cursor.execute("DELETE FROM shipment_custom_fees WHERE shipment_no=?", (shp_no,))

            for f_name, inp_obj in self.custom_fee_inputs.items():
                val = self._safe_float(inp_obj.text)
                if val > 0:
                    custom_total += val
                    cursor.execute(
                        "INSERT INTO shipment_custom_fees (shipment_no, fee_name, amount) VALUES (?,?,?)",
                        (shp_no, f_name, val))

            total_exp = (duty + storage + transport + quality + agri + ins +
                         other_amt + custom_total + fee)
            balance = adv - total_exp

            cursor.execute('''
                INSERT OR REPLACE INTO shipments 
                (shipment_no, client_id, client_name, currency, advance_amount, customs_duty,
                 storage_fee, transport_fee, quality_fee, agriculture_fee, insurance_fee,
                 other_fee_amount, clearance_fee, total_expenses, balance, doc_path,
                 status, created_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ''', (shp_no, client_id, client, curr, adv, duty, storage, transport,
                  quality, agri, ins, other_amt + custom_total, fee, total_exp,
                  balance, doc, status, created_at))
            conn.commit()
            conn.close()

            post_shipment_accounting(shp_no, client, curr, adv, duty, storage, transport,
                                     quality, agri, ins, other_amt + custom_total, fee)

            msg = f"تم الحفظ! إجمالي التكاليف: {fmt_num(total_exp)} | المتبقي: {fmt_num(balance)} {curr}"
            self.lbl_status.text = ar(msg)
            self.lbl_status.color = (0, 1, 0, 1)
        except Exception as e:
            self.lbl_status.text = ar(f"خطأ في الحفظ: {str(e)}")
            self.lbl_status.color = (1, 0, 0, 1)


# ============================================================
# 19. التطبيق الرئيسي
# ============================================================
class CustomsERPApp(App):
    def build(self):
        self.title = "برنامج التخليص الجمركي اليمني - تطوير يحيى القداح"
        sm = ScreenManager()
        sm.add_widget(LoginScreen(name='login'))
        sm.add_widget(HomeScreen(name='home'))
        sm.add_widget(AddShipmentScreen(name='add_shipment'))
        sm.add_widget(ViewShipmentsScreen(name='view_shipments'))
        sm.add_widget(ClientsScreen(name='clients'))
        sm.add_widget(InvoicesScreen(name='invoices'))
        sm.add_widget(CustomerStatementScreen(name='statement'))
        sm.add_widget(ExpensesScreen(name='expenses'))
        sm.add_widget(DashboardScreen(name='dashboard'))
        sm.add_widget(SettingsScreen(name='settings'))
        return sm


if __name__ == '__main__':
    CustomsERPApp().run()